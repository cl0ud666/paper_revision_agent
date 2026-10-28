"""Offline integration fixture: generate, review, export, reopen and render.

Run in the deployment image so the converter and font stack match production.
No external model calls, customer data or service credentials are used.
"""

import argparse
import json
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw

from app.agents.report import write_redline_docx
from app.document.docx_model import DocxSource
from app.document.docx_redline import applied_rounds, export_redline
from app.document.patches import apply_patches, patch_from_model, validate_patch
from app.document.preview import prepare_visual_review, render_pages
from app.document.spans import build_spans
from app.document.structured import formula_xml, preflight


def run(directory):
    directory.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Noto Sans CJK SC"
    style.font.size = Pt(11)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    for name in ("Title", "Heading 1"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].font.underline = False
        for border in list(doc.styles[name].element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.add_heading("论文结构编辑验证", 0)
    doc.add_paragraph("本文件只用于验证修改、取消修改与结构保护，不包含客户论文。")
    doc.add_heading("1方法与格式", 1)
    doc.add_paragraph("目标术语MoCo的字体将单独加粗并调整字号。")
    doc.add_paragraph("新增公式位置：平均值定义为")
    doc.add_paragraph("已有公式保持原样：")._p.append(formula_xml(r"E=mc^2"))
    pic = Image.new("RGB", (180, 40), "#dceadf")
    ImageDraw.Draw(pic).rectangle((10, 10, 170, 30), fill="#467a60")
    image = BytesIO()
    pic.save(image, format="PNG")
    image.seek(0)
    doc.add_picture(image, width=Inches(1.8))
    doc.add_heading("1.1方法定义", 2)
    title = doc.add_paragraph()
    title.add_run("1.2\t")
    title.add_run("国内外研究现状")
    doc.add_heading("2表格与引用", 1)
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, values in zip(
        table.rows, [("方法", "备注"), ("MoCO", "仅统一名称，不修改数据"), ("BERT", "保留原文")]
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.add_paragraph("此处引用动量对比学习方法。")
    doc.add_heading("参考文献", 1)
    evidence = "Kaiming He; Haoqi Fan; Yuxin Wu; Saining Xie; Ross Girshick. Momentum Contrast for Unsupervised Visual Representation Learning. 2020. 10.1109/cvpr42600.2020.00975"
    source_path = directory / "source.docx"
    doc.save(source_path)
    data = source_path.read_bytes()
    original = DocxSource(data).text
    spans = build_spans(original)

    def make(number, needle, action, params, quote=None):
        span = next(s for s in spans if needle in s.text)
        item = {
            "span_id": span.id,
            "action": action,
            "params": params,
            "quote": quote or span.text,
            "issue_id": number,
            "rationale": "按测试方案进行可独立确认的结构修改",
        }
        p = validate_patch(patch_from_model(item, spans=spans, patch_id=f"patch-{number}"), spans)
        assert p.status == "proposed", p.conflict_reason
        preflight(data, p.to_json())
        p.status = "accepted"
        return p

    patches = [
        make(1, "目标术语", "set_font_format", {"font_size_pt": 14, "bold": True}, "MoCo"),
        make(2, "新增公式", "insert_formula", {"latex": r"\bar{x}=\frac{1}{n}\sum_{i=1}^{n}x_i"}),
        make(
            3,
            "| 方法",
            "table_edit",
            {
                "operation": "set_cell_text",
                "row": 1,
                "column": 0,
                "expected_text": "MoCO",
                "text": "MoCo",
            },
        ),
        make(4, "1.2", "set_font_format", {"font_name": "仿宋_GB2312", "font_size_pt": 7.5}),
        make(
            5,
            "本文件只用于验证",
            "insert_paragraphs",
            {
                "paragraphs": [
                    {"text": "待开展对比实验：固定数据划分与评价指标。"},
                    {"text": "实测结果：准确率XXX，召回率XXX（待填，不作有效性结论）。"},
                ]
            },
        ),
    ]
    applied = apply_patches(original, spans, patches)
    assert applied.document == original
    state = {
        "run_id": "structure-qa",
        "visual_review": True,
        "original_document": original,
        "current_document": original,
        "source_docx": str(source_path),
        "applied_patches": [p.to_json() for p in applied.applied],
        "user_request": evidence,
        "audit_findings": [],
    }
    state.update(prepare_visual_review(state, directory))
    assert len(state["final_patches"]) == 5, state.get("review_diagnostics")
    state.update(final_review_confirmed=True, final_review_ready=True, excluded_patch_ids=[])
    output = write_redline_docx(state, run_dir=directory, logger=lambda _: None)
    assert output["redline_summary"]["applied"] == 5, output
    edited = DocxSource(Path(output["docx_path"]).read_bytes())
    assert output["redline_summary"]["comments"] == 5
    Document(output["docx_path"])
    base = DocxSource(data)
    for name, raw in base.parts.items():
        if name.startswith("word/media/") or name == "word/styles.xml":
            assert edited.parts[name] == raw
    old_math = list(
        base.document_root.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
    )[0]
    assert any(
        etree.tostring(n) == etree.tostring(old_math)
        for n in edited.document_root.iter(old_math.tag)
    )
    rendered, cache = render_pages(output["docx_path"], directory / "rendered")
    # Excluding every patch must leave all original content and no new comments.
    state["excluded_patch_ids"] = [p["patch_id"] for p in state["final_patches"]]
    excluded = export_redline(data, applied_rounds(state), str(directory / "excluded.docx"))
    assert not excluded.applied and excluded.comments == 0 and excluded.text == original
    # A failure after confirmation must appear in both final JSON and Word.
    state["excluded_patch_ids"] = []
    for item in state["final_patches"]:
        if item["action"] == "table_edit":
            item["params"]["expected_text"] = "错误的旧值"
    failed = write_redline_docx(state, run_dir=directory / "failure", logger=lambda _: None)
    assert failed["redline_summary"]["skipped"] == 1, failed
    assert any(i["kind"] == "blocked" for i in failed["unfinished_items"])
    result = {
        "pages": rendered["pages"],
        "page_directory": str(cache),
        "review": state["final_patches"],
        "written": output["redline_summary"]["applied"],
        "comments": output["redline_summary"]["comments"],
        "late_failure_in_web_and_word": True,
        "exclude_all_preserves_original": True,
    }
    (directory / "verification.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({k: v for k, v in result.items() if k != "review"}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    run(parser.parse_args().output)
