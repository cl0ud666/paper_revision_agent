"""Offline converter smoke test, run inside the application image.

python /verification/verify_word_formats.py --output /qa
No credentials, model calls, database writes, or uploaded user documents.
"""

import argparse
import json
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from app.document.docx_model import DocxSource, w
from app.document.docx_redline import export_redline
from app.document.spans import build_spans
from app.document.word_formats import normalize_word, validate_word


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="pra-smoke-") as directory:
        root = Path(directory)
        document = Document()
        document.add_heading("Word format verification", 0)
        document.add_paragraph("本文采用MoCo方法进行实验。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Method"
        table.cell(0, 1).text = "Score"
        table.cell(1, 0).text = "Baseline"
        table.cell(1, 1).text = "85.2"
        source = root / "fixture.docx"
        document.save(source)
        results = []
        soffice = shutil.which("soffice")
        for suffix, filter_name in (
            ("doc", "MS Word 97"),
            ("rtf", "Rich Text Format"),
            ("odt", "writer8"),
        ):
            subprocess.run(
                [
                    soffice,
                    f"-env:UserInstallation={(root / ('profile-' + suffix)).as_uri()}",
                    "--headless",
                    "--convert-to",
                    f"{suffix}:{filter_name}",
                    "--outdir",
                    str(root),
                    str(source),
                ],
                check=True,
                timeout=60,
                capture_output=True,
            )
            legacy = (root / f"fixture.{suffix}").read_bytes()
            _, kind = validate_word(legacy, f"fixture.{suffix}")
            assert kind == suffix
            normalized, warnings = normalize_word(legacy, f"fixture.{suffix}")
            assert warnings
            model = DocxSource(normalized)
            assert "本文采用MoCo方法进行实验。" in model.text
            assert "85.2" in model.text
            span = next(s for s in build_spans(model.text) if "MoCo" in s.text)
            start = span.text.index("MoCo")
            patch = {
                "patch_id": "smoke-1",
                "span_label": span.label,
                "before": "MoCo",
                "quote": "MoCo",
                "replace": "MOCO",
                "rationale": "按作者确认的术语写法统一。",
                "start_in_span": start,
                "end_in_span": start + 4,
            }
            target = output / f"{suffix}-revised.docx"
            outcome = export_redline(normalized, [[patch]], str(target))
            assert not outcome.skipped and outcome.comments == 1
            assert "MOCO" in outcome.text and "MoCo" not in outcome.text
            with zipfile.ZipFile(target) as package:
                comments = etree.fromstring(package.read("word/comments.xml"))
                assert len(list(comments.iter(w("comment")))) == 1
                assert "将「MoCo」改为「MOCO」。" in "".join(comments.itertext())
            if suffix == "doc":
                compatible, _ = normalize_word(legacy, "fixture.wps")
                assert "MoCo" in DocxSource(compatible).text
                results.append({"input": "Word OLE-compatible WPS", "status": "passed"})
            results.append({"input": suffix, "status": "passed", "comments": outcome.comments})
        alias, warnings = normalize_word(source.read_bytes(), "fixture.wps")
        assert alias == source.read_bytes() and not warnings
        results.append({"input": "OOXML-compatible WPS", "status": "passed"})
        print(json.dumps(results, ensure_ascii=False))


if __name__ == "__main__":
    main()
