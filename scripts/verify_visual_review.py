"""Offline fixture creation + real LibreOffice/PDF/OOXML integration verification.

--prepare uses the bundled authoring runtime. Verification runs in the app image.
No real papers, credentials, network calls, or live task mutations.
"""

import argparse
import io
import json
import zipfile
from pathlib import Path


def prepare(root):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    from PIL import Image, ImageDraw

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Noto Sans CJK SC"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.header.paragraphs[0].text = "离线验证样稿 · 非用户论文"
    doc.add_heading("对比学习方法研究", 0)
    doc.add_heading("1引言", 1)
    doc.add_paragraph(
        "本文采用MoCo方法学习图像表示，并与已有方法进行比较。术语写法需要统一，实验数据保持不变。"
    )
    doc.add_paragraph(
        "本研究关注方法在小样本场景中的稳定性。以下段落用于检查预览是否保留中文排版、表格和图片，而不是将正文重新拼成网页。"
    )
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for row, values in zip(
        table.rows,
        [("方法", "准确率", "样本数"), ("MOCO", "85.2%", "200"), ("基线", "81.4%", "200")],
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    image = Image.new("RGB", (900, 220), "#edf1ea")
    draw = ImageDraw.Draw(image)
    for x, label in [(40, "Input"), (350, "Encoder"), (660, "Output")]:
        draw.rounded_rectangle((x, 65, x + 200, 155), radius=14, fill="#53694f")
        draw.text((x + 65, 100), label, fill="white", font_size=22)
    draw.line((240, 110, 350, 110), fill="#53694f", width=4)
    draw.line((550, 110, 660, 110), fill="#53694f", width=4)
    png = io.BytesIO()
    image.save(png, format="PNG")
    png.seek(0)
    doc.add_picture(png, width=Inches(5.5))
    doc.add_paragraph("图1方法流程示意")
    doc.add_page_break()
    doc.add_heading("2方法与讨论", 1)
    doc.add_paragraph(
        "本方法能够显著提升模型的表现。该表述需要明确提升的比较对象，不能凭空补充实验结论。"
    )
    doc.add_paragraph("参数FEWSHOT_QUERY_DIR是程序变量名，代码及变量名保持原样，仅可提出建议。")
    doc.add_paragraph("结论：本次样稿验证的是导出与选择一致性，不代表模型能保证学术判断正确。")
    doc.save(root / "source.docx")


def verify(root):
    from app.document.docx_model import DocxSource
    from app.document.docx_redline import export_redline
    from app.document.patches import Patch, apply_patches, decide_patch, validate_patch
    from app.document.preview import prepare_visual_review, render_pages
    from app.document.spans import build_spans

    source = root / "source.docx"
    data = source.read_bytes()
    original = DocxSource(data).text
    current, history = original, []
    for index, (before, after) in enumerate([("MoCo", "MOCO"), ("能够显著提升", "有助于改善")]):
        spans = build_spans(current)
        span = next(s for s in spans if before in s.text)
        item = Patch(
            patch_id=f"p{index}",
            span_id=span.id,
            span_label=span.label,
            quote=before,
            replace=after,
            rationale="统一术语写法。" if index == 0 else "避免未充分论证的绝对化表述。",
        )
        validate_patch(item, spans)
        decide_patch(item, "accept")
        outcome = apply_patches(current, spans, [item])
        assert outcome.applied, item
        current = outcome.document
        history.append({"stage": "apply", "applied": [p.to_json() for p in outcome.applied]})
    state = {
        "original_document": original,
        "current_document": current,
        "revision_history": history,
        "source_docx": str(source),
    }
    result = prepare_visual_review(state, root)
    assert len(result["final_patches"]) == 2, result
    assert len(result["preview_manifest"]["pages"]) == 2, result
    # First rejected, second accepted. The first must remain literally MoCo.
    advice_span = next(s for s in build_spans(original) if "FEWSHOT_QUERY_DIR" in s.text)
    exported = export_redline(
        data,
        [[result["final_patches"][1]]],
        str(root / "revised.docx"),
        suggestions=[
            {
                "kind": "advice",
                "span_label": advice_span.label,
                "suggestion": "请核对变量命名说明与代码是否一致，代码内容未修改。",
            }
        ],
    )
    assert "采用MoCo方法" in exported.text and "有助于改善" in exported.text
    assert "FEWSHOT_QUERY_DIR" in exported.text and len(exported.applied) == 1
    with zipfile.ZipFile(root / "revised.docx") as archive:
        from lxml import etree

        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        comments = etree.fromstring(archive.read("word/comments.xml"))
        texts = ["".join(c.itertext()) for c in comments]
        assert len(texts) == 2 and any(t.startswith("修改建议：") for t in texts), texts
        xml = etree.fromstring(archive.read("word/document.xml"))
        assert len(xml.findall(".//w:ins", ns)) == len(xml.findall(".//w:del", ns)) == 1
    assert source.read_bytes() == data
    render_pages(root / "revised.docx", root / "export-preview")
    (root / "review.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            {"pages": 2, "selectable": 2, "written": 1, "comments": 2, "source_unchanged": True}
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    parser.add_argument("--prepare", action="store_true")
    args = parser.parse_args()
    root = Path(args.output)
    root.mkdir(parents=True, exist_ok=True)
    prepare(root) if args.prepare else verify(root)
