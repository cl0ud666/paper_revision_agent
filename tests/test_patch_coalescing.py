"""Duplicate context windows are not failed edits; choices remain authoritative."""
from itertools import permutations

from app.document.patches import Patch, apply_patches, coalesce_compatible_patches, validate_patch
from app.document.spans import build_spans
from app.document.final_review import compose_changes


def make(spans, identity, quote, replacement, status='accepted'):
    patch = Patch(identity, spans[0].id, quote=quote, replace=replacement, issue_id=identity)
    validate_patch(patch, spans)
    assert patch.status == 'proposed'
    patch.status = status
    return patch


def test_real_moco_duplicate_preserves_text_and_provenance_in_either_order():
    text = '本文对MoCO预训练阶段的超参数进行网格搜索调优。\n'
    spans = build_spans(text)
    patches = [make(spans, 'model', text.strip(), text.strip().replace('MoCO', 'MoCo')),
               make(spans, 'rule', 'MoCO', 'MoCo')]
    for ordered in permutations(patches):
        outcome = apply_patches(text, spans, ordered)
        assert outcome.document == text.replace('MoCO', 'MoCo')
        assert not outcome.skipped
        assert len(outcome.applied) == 1
        final = compose_changes({'original_document': text, 'current_document': outcome.document,
                                 'applied_patches': [p.to_json() for p in outcome.applied]})
        assert len(final) == 1
        assert set(final[0]['source_patch_ids']) == {'model', 'rule'}
        assert set(final[0]['issue_ids']) == {'model', 'rule'}


def test_compatible_changes_in_overlapping_windows_are_both_applied():
    text = '本文使用MoCO，参数为32，结果较好。\n'
    spans = build_spans(text)
    patches = [make(spans, 'a', '本文使用MoCO，参数为32', '本文采用MoCO，参数为32'),
               make(spans, 'b', 'MoCO，参数为32，结果较好', 'MoCo，参数为32，结果很好')]
    for ordered in permutations(patches):
        outcome = apply_patches(text, spans, ordered)
        assert outcome.document == '本文采用MoCo，参数为32，结果很好。\n'
        assert not outcome.skipped


def test_rejected_change_never_joins_accepted_replacement():
    text = '本文使用MoCO。\n'
    spans = build_spans(text)
    patches = [make(spans, 'a', '本文使用MoCO', '本文采用MoCO'),
               make(spans, 'b', 'MoCO', 'MoCo', status='rejected')]
    outcome = apply_patches(text, spans, patches)
    assert outcome.document == '本文采用MoCO。\n'
    assert not outcome.skipped


def test_different_replacements_of_same_word_stay_conflicting():
    text = '使用MoCO。\n'
    spans = build_spans(text)
    patches = [make(spans, 'a', 'MoCO', 'MoCo'), make(spans, 'b', 'MoCO', 'MOCO')]
    outcome = apply_patches(text, spans, patches)
    assert len(outcome.skipped) == 1
    assert outcome.skipped[0].conflict_code == 'overlap'


def test_distinct_positions_with_same_replacement_are_not_deduplicated():
    text = '使用MoCO及MOCO。\n'
    spans = build_spans(text)
    patches = [make(spans, 'a', 'MoCO', 'MoCo'), make(spans, 'b', 'MOCO', 'MoCo')]
    outcome = apply_patches(text, spans, patches)
    assert outcome.document == '使用MoCo及MoCo。\n'
    assert len(outcome.applied) == 2


def test_coalescing_does_not_mutate_input_checkpoint_patches():
    text = '使用MoCO。\n'
    spans = build_spans(text)
    patches = [make(spans, 'a', text.strip(), '使用MoCo。'), make(spans, 'b', 'MoCO', 'MoCo')]
    before = [p.to_json() for p in patches]
    coalesce_compatible_patches(patches, spans)
    assert [p.to_json() for p in patches] == before


def test_merged_patch_exports_once_and_deselecting_restores_original(tmp_path):
    import io
    from docx import Document
    from app.document.docx_model import DocxSource
    from app.document.docx_redline import export_redline
    doc = Document()
    paragraph = doc.add_paragraph('本文对')
    paragraph.add_run('MoCO').bold = True
    paragraph.add_run('预训练阶段的超参数进行网格搜索调优。')
    buffer = io.BytesIO()
    doc.save(buffer)
    source = buffer.getvalue()
    text = DocxSource(source).text
    spans = build_spans(text)
    patches = [make(spans, 'model', spans[0].text, spans[0].text.replace('MoCO', 'MoCo')),
               make(spans, 'rule', 'MoCO', 'MoCo')]
    outcome = apply_patches(text, spans, patches)
    final = compose_changes({'original_document': text, 'current_document': outcome.document,
                             'applied_patches': [p.to_json() for p in outcome.applied]})
    accepted = export_redline(source, [final], str(tmp_path / 'accepted.docx'))
    assert accepted.text == text.replace('MoCO', 'MoCo')
    assert len(accepted.applied) == 1
    assert accepted.comments == 1
    rejected = export_redline(source, [[]], str(tmp_path / 'rejected.docx'))
    assert rejected.text == text
    assert not rejected.applied
    assert rejected.comments == 0
