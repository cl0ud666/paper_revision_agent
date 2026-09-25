"""
文件加载测试：docx / pdf / txt（含 GBK 编码）与错误提示。
"""

from __future__ import annotations

import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from app.utils.loaders import (
    DocumentLoadError,
    decode_text,
    load_bytes,
    load_document,
    supported_extensions,
)

DOC_TEXT = "本文提出了一个新的方法，效果比较好。"


def make_docx(*, stdlib_only: bool = False) -> bytes:
    """生成一个最小可用的 .docx（带标题层级与表格）。"""

    if not stdlib_only:
        try:
            from docx import Document

            document = Document()
            document.add_heading("论文标题", level=1)
            document.add_heading("摘要", level=2)
            document.add_paragraph(DOC_TEXT)
            document.add_heading("1 实验", level=2)
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "方法"
            table.cell(0, 1).text = "准确率"
            table.cell(1, 0).text = "本文方法"
            table.cell(1, 1).text = "85.2"
            buffer = io.BytesIO()
            document.save(buffer)
            return buffer.getvalue()
        except ImportError:
            pass

    # 标准库构造（用于测试内置解析器）
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>论文标题</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Heading2"/></w:pPr><w:r><w:t>摘要</w:t></w:r></w:p>
    <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="ListParagraph"/></w:pPr><w:r><w:t>要点一</w:t></w:r></w:p>
  </w:body>
</w:document>""".format(text=DOC_TEXT)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr("word/document.xml", document_xml)
    return buffer.getvalue()


class TestTextLoading(unittest.TestCase):
    def test_utf8_text(self) -> None:
        loaded = load_bytes(DOC_TEXT.encode("utf-8"), filename="paper.txt")
        self.assertIn("新的方法", loaded.text)
        self.assertEqual(loaded.format, "txt")
        self.assertEqual(loaded.warnings, [])

    def test_gbk_text_with_warning(self) -> None:
        """中文 Windows 记事本存的 ANSI(GBK) 文件也要能读，不能整篇乱码。"""

        loaded = load_bytes(DOC_TEXT.encode("gbk"), filename="paper.txt")
        self.assertIn("新的方法", loaded.text)
        self.assertIn("gb", loaded.encoding.lower())
        self.assertTrue(loaded.warnings)

    def test_utf8_bom(self) -> None:
        loaded = load_bytes("\ufeff测试内容".encode("utf-8"), filename="a.md")
        self.assertNotIn("\ufeff", loaded.text)
        self.assertIn("测试内容", loaded.text)

    def test_decode_text_reports_encoding(self) -> None:
        text, encoding = decode_text("中文".encode("gb18030"))
        self.assertEqual(text, "中文")
        self.assertEqual(encoding, "gb18030")

    def test_empty_file_rejected(self) -> None:
        with self.assertRaises(DocumentLoadError):
            load_bytes(b"", filename="empty.txt")
        with self.assertRaises(DocumentLoadError):
            load_bytes(b"   \n\n  ", filename="blank.txt")

    def test_unsupported_extension_hint(self) -> None:
        with self.assertRaises(DocumentLoadError) as ctx:
            load_bytes(b"fake", filename="old.doc")
        self.assertIn("另存为 .docx", str(ctx.exception))

    def test_load_from_path(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "论文.txt"
            path.write_text(DOC_TEXT, encoding="utf-8")
            loaded = load_document(path)
            self.assertIn("新的方法", loaded.text)
            self.assertEqual(loaded.source, "论文.txt")

    def test_missing_path(self) -> None:
        with self.assertRaises(DocumentLoadError):
            load_document("不存在的文件.docx")

    def test_supported_extensions_text(self) -> None:
        text = supported_extensions()
        for suffix in (".docx", ".pdf", ".md", ".txt"):
            self.assertIn(suffix, text)


class TestDocxLoading(unittest.TestCase):
    def test_omml_formula_is_kept(self) -> None:
        """Word 公式是 OMML 对象：必须取出来，否则模型会误报"公式缺失"。"""

        try:
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError:  # pragma: no cover
            self.skipTest("未安装 python-docx")

        document = Document()
        document.add_heading("3 方法", level=2)
        paragraph = document.add_paragraph("损失函数如下：")

        math_paragraph = document.add_paragraph()
        omath = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        text_node = OxmlElement("m:t")
        text_node.text = "L=-log(exp(s+)/Σexp(s))"
        run.append(text_node)
        omath.append(run)
        math_paragraph._element.append(omath)

        buffer = io.BytesIO()
        document.save(buffer)
        loaded = load_bytes(buffer.getvalue(), filename="formula.docx")

        self.assertIn("损失函数如下", loaded.text)
        self.assertIn("公式", loaded.text)
        self.assertIn("log", loaded.text)

    def test_stdlib_parser_keeps_math(self) -> None:
        from app.utils.loaders import _docx_text_stdlib

        document_xml = (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<w:body><w:p>"
            '<w:r><w:t>损失函数为 </w:t></w:r>'
            "<m:oMath><m:r><m:t>L=-log p</m:t></m:r></m:oMath>"
            "</w:p></w:body></w:document>"
        )
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("word/document.xml", document_xml)
        text = _docx_text_stdlib(buffer.getvalue())
        self.assertIn("损失函数为", text)
        self.assertIn("L=-log p", text)

    def test_docx_via_python_docx(self) -> None:
        loaded = load_bytes(make_docx(), filename="paper.docx")
        self.assertEqual(loaded.format, "docx")
        self.assertIn("# 论文标题", loaded.text)     # 标题层级被保留
        self.assertIn("## 摘要", loaded.text)
        self.assertIn("新的方法", loaded.text)
        self.assertIn("| 方法 | 准确率 |", loaded.text)  # 表格转成 Markdown
        self.assertIn("85.2", loaded.text)

    def test_docx_stdlib_fallback(self) -> None:
        from app.utils.loaders import _docx_text_stdlib

        text = _docx_text_stdlib(make_docx(stdlib_only=True))
        self.assertIn("# 论文标题", text)
        self.assertIn("## 摘要", text)
        self.assertIn("- 要点一", text)

    def test_extracted_docx_becomes_addressable_spans(self) -> None:
        """docx 解析出来的结构，要被段落解析器识别成带 ID 的块。"""

        from app.document.spans import build_spans

        loaded = load_bytes(make_docx(), filename="paper.docx")
        spans = build_spans(loaded.text)
        kinds = {span.kind for span in spans}
        labels = {span.label for span in spans}

        self.assertIn("heading", kinds)
        self.assertIn("paragraph", kinds)
        self.assertIn("T01", labels)  # 表格
        self.assertIn("P01", labels)

    def test_empty_docx_rejected(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("word/document.xml", "<w:document/>")
        buffer.seek(0)
        with self.assertRaises(DocumentLoadError):
            load_bytes(buffer.getvalue(), filename="empty.docx")


class TestPdfLoading(unittest.TestCase):
    def test_blank_pdf_gives_actionable_error(self) -> None:
        """扫描版/空白 PDF 要给出可执行的提示，而不是含糊的崩溃。"""

        try:
            from pypdf import PdfWriter
        except ImportError:  # pragma: no cover - 没装 pypdf 就跳过
            self.skipTest("未安装 pypdf")

        buffer = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.write(buffer)

        with self.assertRaises(DocumentLoadError) as ctx:
            load_bytes(buffer.getvalue(), filename="scan.pdf")
        message = str(ctx.exception)
        self.assertTrue("OCR" in message or "docx" in message)

    def test_pdf_with_text_is_loaded(self) -> None:
        """用打桩的方式验证 PDF 流程接线（真实 PDF 由用户环境解析）。"""

        import app.utils.loaders as loaders

        original = loaders._pdf_text_pypdf
        loaders._pdf_text_pypdf = lambda data: "第 1 页正文。\n\n第 2 页正文。"
        try:
            loaded = load_bytes(b"%PDF-1.4 fake", filename="paper.pdf")
        finally:
            loaders._pdf_text_pypdf = original

        self.assertEqual(loaded.format, "pdf")
        self.assertIn("第 1 页正文", loaded.text)
        self.assertTrue(any("公式" in warning for warning in loaded.warnings))


if __name__ == "__main__":
    unittest.main()
