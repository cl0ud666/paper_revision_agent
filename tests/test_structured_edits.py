"""OOXML patch protocol and immutable-original confirmation regression tests."""

import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch as mock_patch

import pytest
from docx import Document
from lxml import etree

from app.agents import revision
from app.document.docx_model import DocxSource, w
from app.document.docx_redline import applied_rounds, export_redline
from app.document.final_review import compose_changes
from app.document.patches import apply_patches, patch_from_model, validate_patch
from app.document.spans import build_spans
from app.document.structured import M, formula_xml, preflight


def fixture():
    doc = Document()
    doc.add_heading("研究方法", 1)
    p = doc.add_paragraph()
    p.add_run("本文采用 ").bold = True
    p.add_run("MoCo 方法。")
    doc.add_paragraph("损失函数如下。")
    table = doc.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("方法", "准确率"), ("MoCo", "90%"), ("BERT", "80%")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.add_heading("参考文献", 1)
    doc.add_paragraph("作者提供：张三. 机器学习研究. 示例期刊. 2020. 1-10")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def make(data, action, params, needle="本文采用", quote=None):
    source = DocxSource(data)
    spans = build_spans(source.text)
    span = next(s for s in spans if needle in s.text)
    item = {
        "span_id": span.id,
        "span_label": span.label,
        "action": action,
        "params": params,
        "quote": quote or span.text,
        "rationale": "测试用户确认的修改",
        "issue_id": 1,
    }
    p = validate_patch(patch_from_model(item, spans=spans, patch_id="patch-1"), spans)
    assert p.status == "proposed", p.conflict_reason
    return p


def export(tmp_path, data, patch):
    out = tmp_path / "revised.docx"
    result = export_redline(data, [[patch.to_json()]], str(out))
    return result, DocxSource(out.read_bytes())


def test_font_target_runs_only_and_one_comment(tmp_path):
    data = fixture()
    p = make(data, "set_font_format", {"font_size_pt": 14, "font_name": "宋体"}, quote="MoCo")
    preflight(data, p.to_json())
    result, edited = export(tmp_path, data, p)
    assert len(result.applied) == 1 and not result.skipped
    assert result.comments == 1
    assert edited.text == DocxSource(data).text
    runs = edited.document_root.findall(f".//{w('r')}")
    selected = next(r for r in runs if r.findtext(w("t")) == "MoCo")
    assert selected.find(f"{w('rPr')}/{w('sz')}").get(w("val")) == "28"
    assert selected.find(f"{w('rPr')}/{w('rPrChange')}") is not None
    other = next(r for r in runs if r.findtext(w("t")) == "本文采用 ")
    assert other.find(f"{w('rPr')}/{w('b')}") is not None
    assert other.find(f"{w('rPr')}/{w('sz')}") is None
    assert edited.parts["word/styles.xml"] == DocxSource(data).parts["word/styles.xml"]


@pytest.mark.parametrize(
    "params",
    [
        {
            "operation": "set_cell_text",
            "row": 1,
            "column": 0,
            "expected_text": "MoCo",
            "text": "MOCO",
        },
        {"operation": "delete_row", "row": 2},
        {"operation": "delete_column", "column": 1},
        {"operation": "merge_cells", "row": 0, "column": 0, "end_column": 1},
        {
            "operation": "set_table_format",
            "border_color": "999999",
            "shading": "EEEEEE",
            "column_widths_twips": [2200, 2800],
        },
    ],
)
def test_table_operations(tmp_path, params):
    data = fixture()
    p = make(data, "table_edit", params, needle="| 方法")
    preflight(data, p.to_json())
    result, edited = export(tmp_path, data, p)
    assert not result.skipped, result.skipped
    assert result.comments == 1
    Document(BytesIO(Path(result.out_path).read_bytes()))
    if params["operation"] == "set_cell_text":
        assert "MOCO" in edited.text
        assert edited.document_root.find(f".//{w('delText')}").text == "MoCo"
    if params["operation"] == "delete_row":
        assert "BERT" not in edited.text
    original = DocxSource(data)
    assert original.parts["word/styles.xml"] == edited.parts["word/styles.xml"]


def test_split_merged_cell_preserves_content(tmp_path):
    doc = Document(BytesIO(fixture()))
    doc.tables[0].cell(0, 0).merge(doc.tables[0].cell(0, 1))
    stream = BytesIO()
    doc.save(stream)
    data = stream.getvalue()
    p = make(
        data, "table_edit", {"operation": "split_cell", "row": 0, "column": 0}, needle="| 方法"
    )
    result, edited = export(tmp_path, data, p)
    assert not result.skipped
    assert "准确率" in edited.text and "方法" in edited.text
    assert edited.document_root.find(f".//{w('gridSpan')}") is None


def test_formula_is_native_and_preserves_other_nodes(tmp_path):
    data = fixture()
    p = make(
        data,
        "insert_formula",
        {"latex": r"\frac{a}{b}=\sum_{i=1}^{n}x_i", "position": "after"},
        needle="损失函数",
    )
    preflight(data, p.to_json())
    result, edited = export(tmp_path, data, p)
    assert not result.skipped, result.skipped
    math = edited.document_root.find(f".//{{{M}}}oMath")
    assert math is not None
    assert math.getparent().tag == w("ins")
    assert math.find(f".//{{{M}}}f") is not None
    assert len(edited.document_root.findall(f".//{w('tbl')}")) == 1
    assert result.comments == 1


@pytest.mark.parametrize(
    "latex", [r"\input{secret}", r"\href{http://example.com}{x}", r"\unrecognizedcommand{a}"]
)
def test_invalid_formulas_are_not_silently_emitted(latex):
    with pytest.raises(ValueError):
        formula_xml(latex)


@pytest.mark.parametrize(
    "latex",
    [
        r"\bar{x}",
        r"\overline{ab}",
        r"\underline{x}",
        r"\bar{x}=\frac{1}{n}\sum_{i=1}^{n}x_i",
        r"\begin{pmatrix}a&b\\c&d\end{pmatrix}",
    ],
)
def test_accent_fraction_sum_and_matrix_are_well_formed(latex):
    result = formula_xml(latex)
    assert result.tag == f"{{{M}}}oMath"
    assert len(result)


def citation(data, key="ref_zhang2020"):
    source = DocxSource(data)
    bib = next(s for s in build_spans(source.text) if "参考文献" in s.text)
    span = next(s for s in build_spans(source.text) if "本文采用" in s.text)
    return patch_from_model(
        {
            "action": "insert_citation",
            "span_label": span.label,
            "quote": "MoCo",
            "params": {
                "bib_key": key,
                "metadata": {
                    "author": "张三",
                    "title": "机器学习研究",
                    "year": "2020",
                    "journal": "示例期刊",
                    "pages": "1-10",
                },
                "source_evidence": "张三. 机器学习研究. 示例期刊. 2020. 1-10",
                "bibliography_span_label": bib.label,
            },
        },
        spans=build_spans(source.text),
        patch_id="disabled-cite",
    )


def test_disabled_citation_has_no_sources_or_fields_and_has_manual_comment(tmp_path):
    data = fixture()
    p = citation(data)
    with pytest.raises(ValueError, match="已关闭"):
        preflight(data, p.to_json())
    result, edited = export(tmp_path, data, p)
    assert result.skipped and not result.applied
    fields = [n.text for n in edited.document_root.iter(w("instrText"))]
    assert not fields
    original = DocxSource(data)
    assert {k: v for k, v in edited.parts.items() if k.startswith("customXml/")} == {
        k: v for k, v in original.parts.items() if k.startswith("customXml/")
    }
    assert edited.text == DocxSource(data).text
    assert result.comments == 1


@pytest.mark.parametrize(
    "action,params",
    [
        ("insert_table", {}),
        ("table_edit", {"operation": "add_row", "row": 3, "values": ["新方法", "XXX"]}),
        ("table_edit", {"operation": "add_column", "column": 2, "values": ["说明", "甲", "乙"]}),
        ("add_bibliography_item", {}),
    ],
)
def test_disabled_additions_cannot_bypass_validation_or_export(tmp_path, action, params):
    data = fixture()
    spans = build_spans(DocxSource(data).text)
    span = next(s for s in spans if s.kind == "table")
    p = patch_from_model(
        {"span_label": span.label, "quote": span.text, "action": action, "params": params},
        spans=spans,
        patch_id="disabled",
    )
    validate_patch(p, spans)
    assert p.status == "conflict" and "已关闭" in p.conflict_reason
    with pytest.raises(ValueError, match="已关闭"):
        preflight(data, p.to_json())
    result, edited = export(tmp_path, data, p)
    assert not result.applied and result.skipped and result.comments == 1
    assert edited.text == DocxSource(data).text


def test_failed_operation_is_atomic_and_gets_manual_comment(tmp_path):
    data = fixture()
    p = make(
        data,
        "table_edit",
        {
            "operation": "set_cell_text",
            "row": 1,
            "column": 0,
            "expected_text": "错的原文",
            "text": "MOCO",
        },
        needle="| 方法",
    )
    with pytest.raises(ValueError):
        preflight(data, p.to_json())
    result, edited = export(tmp_path, data, p)
    assert not result.applied and len(result.skipped) == 1
    assert edited.text == DocxSource(data).text
    assert result.comments == 1
    assert "无法修改" in "".join(etree.fromstring(edited.parts["word/comments.xml"]).itertext())


def test_structural_patch_flows_through_final_confirmation_and_exclusion(tmp_path):
    data = fixture()
    original = DocxSource(data).text
    p = make(data, "set_font_format", {"bold": True}, quote="MoCo")
    p.status = "accepted"
    result = apply_patches(original, build_spans(original), [p])
    assert result.document == original and len(result.applied) == 1
    state = {
        "original_document": original,
        "current_document": original,
        "applied_patches": [p.to_json()],
        "visual_review": True,
        "final_review_confirmed": True,
    }
    final = compose_changes(state)
    assert len(final) == 1 and final[0]["action"] == "set_font_format"
    state["final_patches"] = final
    state["excluded_patch_ids"] = [final[0]["patch_id"]]
    out = export_redline(data, applied_rounds(state), str(tmp_path / "excluded.docx"))
    assert not out.applied and out.comments == 0
    assert DocxSource((tmp_path / "excluded.docx").read_bytes()).text == original


def test_advice_is_sent_to_revision_instead_of_comment_only():
    document = "本方法效果很好。\n"
    issue = {
        "id": 1,
        "kind": "advice",
        "span_label": "P01",
        "evidence": "本方法效果很好。",
        "problem": "表述过于笼统",
        "suggestion": "改为本方法在本实验中有效。",
    }
    response = SimpleNamespace(
        text=json.dumps(
            {
                "patches": [
                    {
                        "issue_id": 1,
                        "span_label": "P01",
                        "quote": "本方法效果很好。",
                        "replace": "本方法在本实验中有效。",
                        "rationale": "限制结论适用范围",
                    }
                ]
            },
            ensure_ascii=False,
        )
    )
    with mock_patch("app.agents.revision.agent.call_llm", return_value=response) as llm:
        result = revision.run(
            {"current_document": document, "audit_findings": [issue]}, logger=lambda _: None
        )
    assert llm.called and len(result["proposed_patches"]) == 1
    assert not result["manual_items"]


def test_missing_experiment_data_still_generates_placeholder_structure():
    issue = {
        "id": 1,
        "kind": "advice",
        "span_label": "P01",
        "evidence": "需要验证。",
        "problem": "需要验证。",
        "suggestion": "补做消融实验",
    }
    response = SimpleNamespace(
        text=json.dumps(
            {
                "patches": [
                    {
                        "issue_id": 1,
                        "span_label": "P01",
                        "quote": "需要验证。",
                        "action": "insert_paragraphs",
                        "params": {
                            "paragraphs": [
                                {"text": "消融实验设置：待开展实验。"},
                                {"text": "实测准确率：XXX（待填）。"},
                            ]
                        },
                        "rationale": "补充待实施的实验结构，不编造数据",
                    }
                ]
            },
            ensure_ascii=False,
        )
    )
    with mock_patch("app.agents.revision.agent.call_llm", return_value=response) as model:
        result = revision.run(
            {"current_document": "需要验证。\n", "audit_findings": [issue]}, logger=lambda _: None
        )
    assert model.called
    assert not result["manual_items"]
    assert result["proposed_patches"][0]["action"] == "insert_paragraphs"


def test_source_checksum_mismatch_prevents_export(tmp_path):
    data = fixture()
    p = make(data, "set_font_format", {"bold": True}, quote="MoCo").to_json()
    p["original_checksum"] = "0" * 64
    out = export_redline(data, [[p]], str(tmp_path / "mismatch.docx"))
    assert not out.applied and len(out.skipped) == 1
    assert out.comments == 1
    assert DocxSource(Path(out.out_path).read_bytes()).text == DocxSource(data).text


@pytest.mark.parametrize("operation", ["set_cell_text", "delete_row"])
def test_table_operations_protect_code_identifiers(tmp_path, operation):
    doc = Document(BytesIO(fixture()))
    doc.tables[0].cell(1, 0).text = "FEWSHOT_QUERY_DIR"
    stream = BytesIO()
    doc.save(stream)
    data = stream.getvalue()
    params = {"operation": operation, "row": 1}
    if operation == "set_cell_text":
        params.update(column=0, expected_text="FEWSHOT_QUERY_DIR", text="fewshot_query_dir")
    p = make(data, "table_edit", params, needle="| 方法")
    with pytest.raises(ValueError):
        preflight(data, p.to_json())
    result, edited = export(tmp_path, data, p)
    assert not result.applied and len(result.skipped) == 1
    assert edited.text == DocxSource(data).text
