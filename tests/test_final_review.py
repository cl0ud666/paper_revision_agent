"""Frozen original→final plans and conservative page anchoring (no LLM)."""
from pathlib import Path
import tempfile
from unittest.mock import patch
import unittest

from app.document.docx_model import DocxSource
from app.document.final_review import compose_changes, unfinished_items, visible_unfinished
from app.document.preview import locate_on_pages, prepare_visual_review
from app.document.docx_redline import export_redline, applied_rounds
from tests.test_docx_redline import applied_round, _png
from tests.test_web_api import paper_bytes


def fake_pages(source_path, directory):
    """Only the external PDF renderer is stubbed; graph/anchoring/export stay real."""
    text = DocxSource(Path(source_path).read_bytes()).text
    chars = [[c, 1, .1 + (i % 70) * .01, .1 + (i // 70) * .03,
              .109 + (i % 70) * .01, .12 + (i // 70) * .03]
             for i, c in enumerate(text) if not c.isspace()]
    data = {"fingerprint": "test-render", "pages": [{"number": 1, "width": 612, "height": 792}], "chars": chars}
    cache = Path(directory) / data["fingerprint"]
    cache.mkdir(parents=True, exist_ok=True)
    (cache / "page-1.png").write_bytes(_png())
    return data, cache


def state_for(original, replacements):
    text, history = original, []
    for index, (before, after) in enumerate(replacements):
        text, items = applied_round(text, label="P01", quote=before, replace=after)
        items[0]["patch_id"] = f"patch-{index}"
        history.append({"stage": "apply", "applied": items})
    return {"original_document": original, "current_document": text, "revision_history": history}


class TestFinalReview(unittest.TestCase):
    def test_dependent_rounds_become_one_replacement(self):
        state = state_for("本文采用 MoCo 方法。\n", [("MoCo", "MOCO"), ("MOCO", "moco")])
        result = compose_changes(state)
        self.assertEqual([(p["before"], p["replace"]) for p in result], [("MoCo", "moco")])
        self.assertEqual(result[0]["source_patch_ids"], ["patch-0", "patch-1"])

    def test_partial_second_edit_keeps_whole_first_group(self):
        state = state_for("本文采用甲方法。\n", [("甲方法", "优秀乙方法"), ("乙", "丙")])
        self.assertEqual(compose_changes(state)[0]["replace"], "优秀丙方法")

    def test_independent_changes_and_return_to_original(self):
        state = state_for("本文采用 MoCo 方法，表达较差。\n", [("MoCo", "MOCO"), ("较差", "准确")])
        self.assertEqual(len(compose_changes(state)), 2)
        state = state_for("本文采用 MoCo 方法。\n", [("MoCo", "MOCO"), ("MOCO", "MoCo")])
        self.assertEqual(compose_changes(state), [])

    def test_corrupted_history_fails_closed(self):
        state = state_for("本文采用 MoCo 方法。\n", [("MoCo", "MOCO")])
        state["current_document"] += "未记录的改动"
        with self.assertRaises(ValueError):
            compose_changes(state)

    def test_export_subset_and_no_double_insertion(self):
        data = paper_bytes()
        state = state_for(DocxSource(data).text, [("MoCo", "moco"), ("moco", "MOCO"), ("实验", "研究")])
        final = compose_changes(state)
        state.update(visual_review=True, final_review_ready=True, final_review_confirmed=True,
                     final_patches=final, excluded_patch_ids=[final[0]["patch_id"]])
        with tempfile.TemporaryDirectory() as directory:
            outcome = export_redline(data, applied_rounds(state), str(Path(directory) / "result.docx"))
            self.assertIn("MoCo 方法进行研究", outcome.text)
            self.assertNotIn("MoCoMOCO", outcome.text)
            self.assertEqual(len(outcome.applied), 1)
            self.assertEqual(outcome.comments, 1)

    def test_unconfirmed_export_is_forbidden(self):
        with self.assertRaises(ValueError):
            applied_rounds({"visual_review": True, "final_review_ready": True})

    def test_paragraph_deletion_does_not_shift_followup_or_advice(self):
        import io
        from docx import Document
        from lxml import etree
        import zipfile
        doc = Document()
        doc.add_paragraph("第一段需要删除。")
        doc.add_paragraph("第二段需要保留。")
        doc.add_paragraph("第三段提供建议。")
        raw = io.BytesIO(); doc.save(raw); data = raw.getvalue()
        original = DocxSource(data).text
        updated, first = applied_round(original, label="P01", quote="第一段需要删除。", replace="", delete=True)
        # The remaining paragraph has a new readable label after deletion.
        updated, second = applied_round(updated, label="P01", quote="保留", replace="保存")
        state = {"original_document": original, "current_document": updated,
                 "revision_history": [{"stage": "apply", "applied": first}, {"stage": "apply", "applied": second}]}
        final = compose_changes(state)
        self.assertEqual(len(final), 2)
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "result.docx"
            result = export_redline(data, [final], str(target), suggestions=[{
                "kind": "advice", "span_label": "P03", "original_anchor_start": original.index("第三段"), "suggestion": "请核对说明。"}])
            self.assertEqual(result.text, "第二段需要保存。\n\n第三段提供建议。\n")
            with zipfile.ZipFile(target) as archive:
                xml = etree.fromstring(archive.read("word/document.xml"))
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                paragraph = next(p for p in xml.findall(".//w:p", ns) if "第三段" in "".join(p.itertext()))
                self.assertIsNotNone(paragraph.find("w:commentRangeStart", ns))

    def test_ambiguous_identical_paragraphs_are_blocked(self):
        original = "本文 MoCo 方法。\n本文 MoCo 方法。\n"
        item = {"original_start": 3, "original_end": 7}
        chars = [[c, 1, .1, .1, .12, .12] for c in original if not c.isspace()]
        with self.assertRaises(ValueError):
            locate_on_pages(item, original, {"chars": chars})

    def test_repeated_token_uses_context(self):
        original = "本文 MoCo 方法优于另一 MoCo 方法。\n"
        item = {"original_start": 3, "original_end": 7}
        chars = [[c, 1, .1 + i * .01, .1, .109 + i * .01, .12] for i, c in enumerate(original) if not c.isspace()]
        self.assertTrue(locate_on_pages(item, original, {"chars": chars}))

    def test_renderer_failure_blocks_changes_and_becomes_advice(self):
        data = paper_bytes()
        state = state_for(DocxSource(data).text, [("MoCo", "MOCO")])
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            source.write_bytes(data)
            state["source_docx"] = str(source)
            with patch("app.document.preview.render_pages", side_effect=ValueError("转换器不可用")):
                result = prepare_visual_review(state, directory)
        self.assertEqual(result["final_patches"], [])
        self.assertEqual(result["unfinished_items"], [])
        self.assertTrue(result["review_diagnostics"])
        self.assertTrue(result["final_manual_items"])

    def test_budget_failed_audit_and_unapplied_issue_reported(self):
        items = unfinished_items({"audit_findings": [{"id": 1, "problem": "未修复的表达"}], "audit_meta": {"failed_chunks": 1}})
        self.assertEqual(len(items), 2)

    def test_visual_unfinished_only_includes_reviewed_modifications(self):
        item = {"patch_id": "final-1", "before": "原文", "span_label": "P01", "issue_ids": [1]}
        state = {"visual_review": True, "final_patches": [item],
                 "audit_findings": [{"id": 2, "problem": "GNN 应为 KNN"}],
                 "conflict_patches": [{"patch_id": "never-shown", "before": "GNN"}],
                 "review_result": {"per_issue": [{"id": 1, "status": "unresolved", "reason": "仍未解决歧义"}],
                                   "remaining_issues": [{"id": 2, "problem": "GNN 只出现一次"}]}}
        result = unfinished_items(state)
        self.assertEqual(result, [])
        self.assertNotIn("GNN", str(result))
        self.assertEqual(unfinished_items({**state, "excluded_patch_ids": ["final-1"]}), [])

    def test_saved_legacy_summary_cannot_leak_unreviewed_candidates(self):
        old = [{"title": "GNN 应为 KNN", "kind": "unresolved"},
               {"title": "未展示的补丁", "kind": "blocked"},
               {"title": "用户取消", "kind": "not_selected"}]
        self.assertEqual(visible_unfinished(old), [old[-1]])
        self.assertEqual(visible_unfinished(old, []), [])
        failed = {"patch_id": "final-1", "kind": "blocked", "title": "已确认但写回失败"}
        self.assertEqual(visible_unfinished(old + [failed], ["final-1"]), [failed])

    def test_old_report_section_is_filtered_without_changing_other_sections(self):
        from app.document.final_review import refresh_unfinished_section
        text = "# 报告\n\n## 审查记录\n历史记录\n\n## 未完成的修改及原因\nGNN 应为 KNN\n\n## 异常与降级记录\n无\n"
        updated = refresh_unfinished_section(text, [])
        self.assertNotIn("GNN 应为 KNN", updated)
        self.assertIn("## 审查记录\n历史记录", updated)
        self.assertTrue(updated.endswith("## 异常与降级记录\n无\n"))

    def test_gate_resume_drops_old_unreviewed_rows(self):
        from app.workflow.nodes import _confirm_before_writing
        state = {**state_for("本文采用 MoCo 方法。\n", [("MoCo", "MOCO")]),
                 "visual_review": True, "final_review_ready": True, "run_id": "test-final",
                 "unfinished_items": [{"title": "GNN 应为 KNN", "kind": "unresolved"}]}
        state["final_patches"] = compose_changes(state)
        with patch("app.workflow.nodes.interrupt", return_value={"excluded_ids": []}):
            result = _confirm_before_writing(state, lambda _: None)
        self.assertEqual(result["unfinished_items"], [])

    def test_new_decision_after_first_gate_becomes_advice(self):
        from app.workflow.nodes import audit_node, issue_review_node, patch_review_node
        from app.config import reset_settings
        import os
        state = {"visual_review": True, "decisions": True, "run_id": "test-final", "audit_findings": []}
        with patch("app.workflow.nodes.audit_agent.run", return_value={"audit_findings": [{"id": 20, "kind": "decision", "problem": "新的歧义"}]}):
            update = audit_node(state, logger=lambda _: None)
        self.assertEqual(update["audit_findings"][0]["kind"], "advice")
        self.assertTrue(update["deferred_advice"])
        with patch("app.workflow.nodes.interrupt", side_effect=AssertionError("no extra gate")), patch.dict(os.environ, {"PRA_PATCH_GATE": "1"}):
            reset_settings()
            issue_review_node({**state, **update}, logger=lambda _: None)
            patch_review_node({**state, "proposed_patches": []}, logger=lambda _: None)
        reset_settings()
