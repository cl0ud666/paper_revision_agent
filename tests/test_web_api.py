"""Web业务与真实LangGraph/DOCX测试；模型和Redis限流使用离线替身。"""

import io
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.config import reset_settings
from app.utils.llm import reset_llm_factory
from app.utils.resilience import reset_guards
from app.utils.tracing import reset_recorders
from app.web import service
from app.web.api import app
from app.web.db import reset_engines, transaction
from app.web.demo import install_demo
from app.web.dispatcher import dispatch_once
from app.web.models import Job, Run
from app.web.settings import settings
from app.web.storage import store
from app.web.worker import execute_job
from app.workflow.graph import reset_checkpointers


def paper_bytes():
    from docx import Document

    document = Document()
    document.add_heading("论文测试", 0)
    document.add_heading("1 引言", 1)
    document.add_paragraph("本文使用 MoCo 方法进行实验，表格中将其写作 MOCO。")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class WebTestCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="pra-web-")
        root = Path(self.temp.name)
        self.environment = patch.dict(
            os.environ,
            {
                "PRA_WEB_DATABASE_URL": f"sqlite:///{root / 'web.sqlite'}",
                "PRA_STORAGE_BACKEND": "local",
                "PRA_LOCAL_STORAGE_ROOT": str(root / "files"),
                "PRA_DATA_DIR": str(root),
                "PRA_RUNS_DIR": str(root / "runs"),
                "PRA_CHECKPOINT_BACKEND": "sqlite",
                "PRA_CHECKPOINT_DB": str(root / "checkpoint.sqlite"),
                "PRA_CHECKPOINT_URL": "",
                "PRA_PUBLIC_ORIGIN": "http://testserver",
                "PRA_RETRY_MAX_ATTEMPTS": "1",
                "PRA_PATCH_GATE": "0",
                "PRA_FINAL_CONFIRM": "1",
                "PRA_DEFAULT_DAILY_RUNS": "5",
                "PRA_DEFAULT_RUN_TOKENS": "400000",
                "PRA_FREE_RUNS": "3",
                "PRA_REGISTRATION_OPEN": "1",
                "PRA_BUDGET_RUN_TOKENS": "0",
                "PRA_BUDGET_TURN_TOKENS": "0",
            },
        )
        self.environment.start()
        settings.cache_clear()
        store.cache_clear()
        reset_settings()
        reset_engines()
        reset_checkpointers()
        reset_guards()
        reset_recorders()
        command.upgrade(Config("alembic.ini"), "head")
        service.add_user("admin", "testing-password-123", admin=True)
        service.add_user("member", "testing-password-123")
        self.client = TestClient(app)
        self.rate = patch("app.web.routes.auth.login_rate_limit")
        self.rate.start()
        self.slot = patch("app.web.worker.acquire_llm_slot", return_value="test-slot")
        self.slot.start()
        self.release = patch("app.web.worker.release_llm_slot")
        self.release.start()
        self.login()
        install_demo()
        from tests.test_final_review import fake_pages

        self.preview = patch("app.document.preview.render_pages", side_effect=fake_pages)
        self.preview.start()

    def tearDown(self):
        self.client.close()
        self.rate.stop()
        self.slot.stop()
        self.release.stop()
        self.preview.stop()
        reset_llm_factory()
        reset_guards()
        reset_recorders()
        reset_checkpointers()
        reset_engines()
        settings.cache_clear()
        store.cache_clear()
        self.environment.stop()
        reset_settings()
        self.temp.cleanup()

    def login(self, username="admin"):
        response = self.client.post(
            "/api/v1/auth/login", json={"username": username, "password": "testing-password-123"}
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.client.headers["X-CSRF-Token"] = response.json()["csrf_token"]
        return response.json()

    def upload(self):
        response = self.client.post(
            "/api/v1/runs", files={"file": ("术语论文.docx", paper_bytes())}
        )
        self.assertEqual(response.status_code, 202, response.text)
        return response.json()

    def run_job(self, run_id):
        with transaction() as db:
            job = db.scalar(select(Job).where(Job.run_id == run_id).order_by(Job.version.desc()))
            job_id = job.id
        execute_job(job_id)
        return self.client.get(f"/api/v1/runs/{run_id}").json(), job_id

    def answer(self, run, payload):
        return self.client.post(
            f"/api/v1/runs/{run['id']}/decisions",
            json={
                "pending_id": run["pending"]["id"],
                "version": run["version"],
                "payload": payload,
            },
        )


class TestWebLifecycle(WebTestCase):
    def test_recent_events_do_not_freeze_after_first_hundred(self):
        run = self.upload()
        with transaction() as db:
            row = db.get(Run, run["id"])
            for index in range(125):
                service.event(db, row, "progress", f"event-{index}")
        url = f"/api/v1/runs/{run['id']}/events"
        first = self.client.get(url).json()
        self.assertEqual(len(first["events"]), 100)
        tail = self.client.get(url, params={"latest": True}).json()
        self.assertEqual(len(tail["events"]), 100)
        self.assertEqual(tail["events"][-1]["message"], "event-124")
        self.assertGreater(tail["cursor"], first["cursor"])
        next_page = self.client.get(url, params={"cursor": first["cursor"]}).json()
        self.assertEqual(next_page["events"][-1]["message"], "event-124")
        self.login("member")
        self.assertEqual(self.client.get(url, params={"latest": True}).status_code, 404)

    def test_selected_manual_option_survives_restart_and_becomes_comment(self):
        from app.document.docx_model import DocxSource
        from app.document.spans import build_spans

        original = DocxSource(paper_bytes()).text
        span = next(s for s in build_spans(original) if s.kind == "paragraph")
        issue = {
            "id": 7,
            "kind": "decision",
            "span_id": span.id,
            "span_label": span.label,
            "problem": "对比实验需要作者确认",
            "evidence": span.text,
            "options": ["删除相关实验描述", "补充 BERT 对比实验，保留描述"],
            "suggestion": "请作者选择",
        }
        run = self.upload()
        with patch(
            "app.agents.audit.run", return_value={"audit_findings": [issue], "issue_counter": 7}
        ):
            run, _ = self.run_job(run["id"])
        self.assertEqual(run["pending"]["request"]["type"], "issue_review")
        self.assertEqual(self.answer(run, {"choices": {"7": 1}}).status_code, 202)
        reset_checkpointers()
        reset_recorders()
        run, _ = self.run_job(run["id"])
        request = run["pending"]["request"]
        self.assertEqual(request["type"], "final_confirm")
        self.assertEqual(request["patches"], [])
        self.assertEqual(len(request["manual_items"]), 1)
        self.assertIn("补充 BERT", request["manual_items"][0]["suggestion"])
        self.assertNotIn("P01", request["manual_items"][0]["position"])
        reset_checkpointers()
        reset_recorders()
        self.assertEqual(self.answer(run, {"excluded_ids": []}).status_code, 202)
        run, _ = self.run_job(run["id"])
        self.assertEqual(run["status"], "completed", run)
        self.assertEqual(len(run["summary"]["unfinished"]), 1)
        output = self.client.get(f"/api/v1/runs/{run['id']}/download").content
        self.assertEqual(DocxSource(output).text, original)
        with zipfile.ZipFile(io.BytesIO(output)) as archive:
            note = archive.read("word/comments.xml").decode()
        self.assertIn("补充 BERT 对比实验，保留描述", note)
        self.assertNotIn("删除相关实验描述", note)

    def test_old_result_and_report_do_not_show_unreviewed_gnn(self):
        run = self.upload()
        with transaction() as db:
            saved = db.get(Run, run["id"])
            saved.status = "completed"
            saved.state_summary = {
                "unfinished": [
                    {"title": "GNN 应为 KNN", "kind": "unresolved"},
                    {"title": "已取消项", "kind": "not_selected", "reason": "用户取消"},
                ]
            }
            saved.report_key = f"test-reports/{saved.id}.md"
            store().put(
                saved.report_key,
                "# 报告\n\n## 未完成的修改及原因\nGNN 应为 KNN\n".encode(),
                "text/markdown",
            )
        detail = self.client.get(f"/api/v1/runs/{run['id']}").json()
        self.assertEqual([i["title"] for i in detail["summary"]["unfinished"]], ["已取消项"])
        report = self.client.get(f"/api/v1/runs/{run['id']}/download?kind=report").text
        self.assertNotIn("GNN 应为 KNN", report)
        self.assertIn("已取消项", report)

    def test_visual_gate_is_frozen_private_and_no_model_calls_after_confirm(self):
        run = self.upload()
        run, _ = self.run_job(run["id"])
        issue = run["pending"]["request"]["issues"][0]
        self.answer(run, {"choices": {str(issue["id"]): 0}})
        run, _ = self.run_job(run["id"])
        request = run["pending"]["request"]
        self.assertTrue(request["visual"])
        self.assertTrue(request["patches"][0]["rects"])
        image_url = (
            f"/api/v1/runs/{run['id']}/preview/1?fingerprint={request['preview']['fingerprint']}"
        )
        self.assertEqual(self.client.get(image_url).status_code, 200)
        self.assertEqual(
            self.client.get(image_url.replace("/preview/1?", "/preview/2?")).status_code, 404
        )
        self.assertEqual(self.client.get(image_url + "invalid").status_code, 404)
        self.login("member")
        self.assertEqual(self.client.get(image_url).status_code, 404)
        self.login()
        reset_checkpointers()
        reset_recorders()
        again = self.client.get(f"/api/v1/runs/{run['id']}").json()
        self.assertEqual(again["pending"]["request"], request)
        self.assertEqual(self.answer(run, {"excluded_ids": ["unknown"]}).status_code, 400)
        self.assertEqual(self.answer(run, {"exclude_all": True}).status_code, 202)
        with (
            patch("app.agents.audit.run", side_effect=AssertionError("no audit after final")),
            patch("app.agents.revision.run", side_effect=AssertionError("no revision after final")),
            patch("app.agents.review.run", side_effect=AssertionError("no review after final")),
            patch(
                "app.document.preview.prepare_visual_review",
                side_effect=AssertionError("frozen plan"),
            ),
        ):
            run, _ = self.run_job(run["id"])
        self.assertEqual(run["status"], "completed", run)
        self.assertEqual(run["summary"]["written"], 0)
        self.assertTrue(any(i["kind"] == "not_selected" for i in run["summary"]["unfinished"]))
        with transaction() as db:
            source = store().get(db.get(Run, run["id"]).source_key)
        self.assertEqual(self.client.get(f"/api/v1/runs/{run['id']}/download").content, source)

    def test_legacy_upload_normalizes_once_and_exclude_all_returns_docx(self):
        original = b"{\\rtf1 input paper}"
        normalized = paper_bytes()
        response = self.client.post("/api/v1/runs", files={"file": ("论文.rtf", original)})
        self.assertEqual(response.status_code, 202, response.text)
        run = response.json()
        with patch(
            "app.web.worker.normalize_word", return_value=(normalized, ["converted"])
        ) as convert:
            run, _ = self.run_job(run["id"])
            self.assertEqual(run["status"], "waiting_user", run)
            issue = run["pending"]["request"]["issues"][0]
            self.answer(run, {"choices": {str(issue["id"]): 0}})
            reset_checkpointers()
            run, _ = self.run_job(run["id"])
            self.answer(run, {"exclude_all": True})
            run, _ = self.run_job(run["id"])
            self.assertEqual(run["status"], "completed", run)
            convert.assert_called_once_with(original, "论文.rtf")
        with transaction() as db:
            self.assertEqual(store().get(db.get(Run, run["id"]).source_key), original)
        downloaded = self.client.get(f"/api/v1/runs/{run['id']}/download")
        self.assertEqual(downloaded.content, normalized)
        self.assertNotIn(".rtf", downloaded.headers["content-disposition"])
        self.assertTrue(run["summary"]["warnings"])

    def test_conversion_failure_is_shown_before_any_model_call(self):
        from app.document.word_formats import WordFormatError

        response = self.client.post("/api/v1/runs", files={"file": ("论文.rtf", b"{\\rtf1 hello}")})
        self.assertEqual(response.status_code, 202)
        with patch(
            "app.web.worker.normalize_word",
            side_effect=WordFormatError("服务器未安装 LibreOffice 转换器"),
        ):
            run, _ = self.run_job(response.json()["id"])
        self.assertEqual(run["status"], "failed")
        self.assertIn("LibreOffice", run["error"])
        self.assertNotIn("used_tokens", run)
        with transaction() as db:
            self.assertEqual(db.get(Run, run["id"]).used_tokens, 0)

    def test_upload_choose_resume_confirm_download_and_duplicate_delivery(self):
        run = self.upload()
        run, first_job = self.run_job(run["id"])
        self.assertEqual(run["status"], "waiting_user", run)
        self.assertEqual(run["pending"]["request"]["type"], "issue_review")
        issue = run["pending"]["request"]["issues"][0]
        selected = next(i for i, option in enumerate(issue["options"]) if "MOCO" in option)
        response = self.answer(run, {"choices": {str(issue["id"]): selected}})
        self.assertEqual(response.status_code, 202, response.text)
        # 双击提交同一个回答不会投递第二个恢复任务。
        self.assertEqual(
            self.answer(run, {"choices": {str(issue["id"]): selected}}).status_code, 202
        )
        self.assertEqual(
            self.answer(run, {"choices": {str(issue["id"]): (selected + 1) % 2}}).status_code, 409
        )
        reset_checkpointers()
        reset_recorders()
        run, _ = self.run_job(run["id"])
        self.assertEqual(run["pending"]["request"]["type"], "final_confirm", run)
        self.assertEqual(self.answer(run, {"excluded_ids": []}).status_code, 202)
        run, final_job = self.run_job(run["id"])
        self.assertEqual(run["status"], "completed", run)
        self.assertNotIn("used_tokens", run)
        with transaction() as db:
            used_tokens = db.get(Run, run["id"]).used_tokens
            self.assertGreater(used_tokens, 0)
        data = self.client.get(f"/api/v1/runs/{run['id']}/download")
        self.assertEqual(data.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(data.content)) as archive:
            xml = archive.read("word/document.xml")
        self.assertIn(b"<w:ins", xml)
        self.assertIn(b"MOCO", xml)
        self.assertEqual(
            self.client.get(f"/api/v1/runs/{run['id']}/download?kind=report").status_code, 200
        )
        execute_job(first_job)
        execute_job(final_job)
        current = self.client.get(f"/api/v1/runs/{run['id']}").json()
        self.assertNotIn("used_tokens", current)
        with transaction() as db:
            self.assertEqual(db.get(Run, run["id"]).used_tokens, used_tokens)

    def test_final_exclude_all_still_returns_original_word_and_report(self):
        run = self.upload()
        run, _ = self.run_job(run["id"])
        issue = run["pending"]["request"]["issues"][0]
        self.answer(run, {"choices": {str(issue["id"]): 0}})
        run, _ = self.run_job(run["id"])
        self.answer(run, {"exclude_all": True})
        run, _ = self.run_job(run["id"])
        self.assertEqual(run["status"], "completed", run)
        with transaction() as db:
            record = db.get(Run, run["id"])
            original = store().get(record.source_key)
        self.assertEqual(self.client.get(f"/api/v1/runs/{run['id']}/download").content, original)

    def test_checkpoint_ahead_of_business_state_does_not_answer_next_gate_twice(self):
        run = self.upload()
        run, _ = self.run_job(run["id"])
        issue = run["pending"]["request"]["issues"][0]
        self.answer(run, {"choices": {str(issue["id"]): 0}})
        from app.web.worker import event as original_event

        def fail_after_checkpoint(db, row, kind, message):
            if kind == "waiting_user":
                raise RuntimeError("simulated crash after checkpoint")
            return original_event(db, row, kind, message)

        with patch("app.web.worker.event", side_effect=fail_after_checkpoint):
            run, _ = self.run_job(run["id"])
        self.assertEqual(run["status"], "failed")
        self.assertEqual(self.client.post(f"/api/v1/runs/{run['id']}/retry").status_code, 202)
        reset_checkpointers()
        run, _ = self.run_job(run["id"])
        self.assertEqual(run["status"], "waiting_user")
        self.assertEqual(run["pending"]["request"]["type"], "final_confirm")

    def test_outbox_survives_broker_failure(self):
        run = self.upload()
        with patch(
            "app.web.dispatcher.celery_app.send_task", side_effect=ConnectionError("offline")
        ):
            self.assertEqual(dispatch_once(), 0)
        with transaction() as db:
            job = db.scalar(select(Job).where(Job.run_id == run["id"]))
            self.assertEqual(job.status, "pending")
            from app.web.models import now

            job.publish_after = now()
        published = []
        self.assertEqual(dispatch_once(published.append), 1)
        self.assertEqual(len(published), 1)


class TestWebAccess(WebTestCase):
    def test_user_cannot_access_another_users_task_choices_events_or_download(self):
        run = self.upload()
        run, _ = self.run_job(run["id"])
        self.login("member")
        for suffix in ("", "/events", "/download"):
            self.assertEqual(self.client.get(f"/api/v1/runs/{run['id']}{suffix}").status_code, 404)
        self.assertEqual(self.answer(run, {"ignore_all": True}).status_code, 404)
        self.assertEqual(self.client.post(f"/api/v1/runs/{run['id']}/cancel").status_code, 404)
        self.assertEqual(self.client.get("/api/v1/admin/users").status_code, 403)

    def test_csrf_expired_session_and_cross_origin_requests(self):
        self.client.headers.pop("X-CSRF-Token")
        self.assertEqual(
            self.client.post(
                "/api/v1/runs", files={"file": ("test.docx", paper_bytes())}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.post(
                "/api/v1/auth/login",
                json={"username": "admin", "password": "testing-password-123"},
                headers={"Origin": "https://attacker.invalid"},
            ).status_code,
            403,
        )
        self.client.cookies.clear()
        self.assertEqual(self.client.get("/api/v1/runs").status_code, 401)

    def test_invite_quota_disable_and_login(self):
        new = self.client.post(
            "/api/v1/admin/users",
            json={
                "username": "newuser",
                "password": "testing-password-123",
                "daily_run_limit": 1,
                "run_token_limit": 123000,
            },
        )
        self.assertEqual(new.status_code, 201, new.text)
        user = new.json()
        self.login("newuser")
        self.upload()
        response = self.client.post("/api/v1/runs", files={"file": ("second.docx", paper_bytes())})
        self.assertEqual(response.status_code, 429)
        self.login()
        self.assertEqual(
            self.client.patch(
                f"/api/v1/admin/users/{user['id']}", json={"active": False}
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.post(
                "/api/v1/auth/login",
                json={"username": "newuser", "password": "testing-password-123"},
            ).status_code,
            401,
        )

    def test_cancelled_queued_job_does_not_call_workflow(self):
        run = self.upload()
        self.assertEqual(
            self.client.post(f"/api/v1/runs/{run['id']}/cancel").json()["status"], "cancelled"
        )
        with patch("app.web.worker.drive_workflow") as drive:
            self.run_job(run["id"])
            drive.assert_not_called()

    def test_repeated_cancel_during_execution_cannot_be_overwritten_by_result(self):
        from app.workflow.harness import RunResult

        run = self.upload()

        def cancel_while_running(record, job, harness):
            for _ in range(2):
                response = self.client.post(f"/api/v1/runs/{run['id']}/cancel")
                self.assertEqual(response.json()["status"], "cancel_requested")
            return RunResult(record.id, harness.thread_id, "completed", state={})

        with patch("app.web.worker.drive_workflow", side_effect=cancel_while_running):
            result, _ = self.run_job(run["id"])
        self.assertEqual(result["status"], "cancelled")
        self.assertFalse(result["has_result"])
        self.assertEqual(self.client.get(f"/api/v1/runs/{run['id']}/download").status_code, 409)

    def test_password_reset_revokes_sessions_and_old_password(self):
        service.reset_password("admin", "replacement-password-456")
        self.assertEqual(self.client.get("/api/v1/users/me").status_code, 401)
        self.assertEqual(
            self.client.post(
                "/api/v1/auth/login", json={"username": "admin", "password": "testing-password-123"}
            ).status_code,
            401,
        )
        self.assertEqual(
            self.client.post(
                "/api/v1/auth/login",
                json={"username": "admin", "password": "replacement-password-456"},
            ).status_code,
            200,
        )

    def test_invalid_files_options_and_stale_decision(self):
        response = self.client.post("/api/v1/runs", files={"file": ("paper.pdf", b"%PDF-test")})
        self.assertEqual(response.status_code, 400)
        run = self.upload()
        run, _ = self.run_job(run["id"])
        issue = run["pending"]["request"]["issues"][0]
        self.assertEqual(self.answer(run, {"choices": {str(issue["id"]): 999}}).status_code, 400)
        response = self.client.post(
            f"/api/v1/runs/{run['id']}/decisions",
            json={"pending_id": run["pending"]["id"], "version": 99, "payload": {}},
        )
        self.assertEqual(response.status_code, 409)
