"""
Word 回写测试：文字与 XML 的锚点、Word 原生修订痕迹、批注，以及各种边界。

这里刻意把论文里最容易出问题的东西放进同一份文档：
标题层级、公式（OMML）、内嵌图片、表格、带换行的代码段。
要验证的不是"能跑通"，而是：

- 读出来的文字与加载器**逐字一致**（否则字符偏移就全错了）；
- 改动只落在该落的地方，其余内容一个字都不动；
- 落不下去的改动被**拒绝并说明原因**，而不是猜着改。
"""

from __future__ import annotations

import io
import struct
import tempfile
import unittest
import zipfile
import zlib
from pathlib import Path

from app.document.docx_model import DocxSource
from app.document.patches import (
    STATUS_APPLIED,
    STATUS_PROPOSED,
    Patch,
    apply_patches,
    decide_patch,
    validate_patch,
)
from app.document.spans import build_spans
from app.utils.loaders import load_bytes

try:  # python-docx / lxml 都是可选依赖
    from docx import Document
    from docx.oxml import OxmlElement

    HAS_DOCX = True
except ImportError:  # pragma: no cover - 取决于运行环境
    HAS_DOCX = False


BODY = "本文提出的方法在三个公开数据集上取得了更好的效果。"
SECOND = "为了验证有效性，我们进行了充分的消融实验。"
CODE = "for i in range(n):\n    total += weight[i]"
FORMULA_TAIL = "定义为 "


def _png() -> bytes:
    """一张 2x2 的 PNG（用来测试"段落里的图片"）。"""

    size = 2
    raw = b"".join(b"\x00" + b"\xff\x40\x00" * size for _ in range(size))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def build_paper() -> bytes:
    """构造一份"五脏俱全"的论文：标题、正文、公式、图片、代码、表格。"""

    document = Document()
    document.add_heading("面向小样本场景的图神经网络方法", level=1)
    document.add_heading("摘要", level=2)
    document.add_paragraph(BODY)
    document.add_heading("3 方法", level=2)
    document.add_paragraph(SECOND)

    formula_paragraph = document.add_paragraph("损失函数" + FORMULA_TAIL)
    omath = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = "L=-log(exp(s+)/Σexp(s))"
    math_run.append(math_text)
    omath.append(math_run)
    formula_paragraph._element.append(omath)

    picture_paragraph = document.add_paragraph("图 3.1 模型结构：")
    picture_paragraph.add_run().add_picture(io.BytesIO(_png()))

    code_paragraph = document.add_paragraph()
    for index, line in enumerate(CODE.split("\n")):
        if index:
            code_paragraph.add_run().add_break()
        code_paragraph.add_run(line)

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "方法"
    table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = "本文方法"
    table.cell(1, 1).text = "85.2"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def applied_round(text: str, *, label: str, quote: str, replace: str, **extra):
    """走一遍真实流程：校验 -> 用户接受 -> 应用，返回（新文字, 本轮已落实的补丁）。"""

    spans = build_spans(text)
    patch = Patch(
        patch_id="patch-001",
        span_id="",
        span_label=label,
        quote=quote,
        replace=replace,
        rationale="让表述更具体。",
        category=extra.pop("category", "语言表达"),
        severity=extra.pop("severity", "high"),
        **extra,
    )
    patch = validate_patch(patch, spans)
    assert patch.status == STATUS_PROPOSED, patch.conflict_reason
    decide_patch(patch, "accept")
    outcome = apply_patches(text, spans, [patch])
    assert outcome.applied, "补丁没有被应用"
    assert outcome.applied[0].status == STATUS_APPLIED
    return outcome.document, [item.to_json() for item in outcome.applied]


def export(data: bytes, rounds, **kwargs):
    """回写到临时目录里的副本，返回 (结果, 文件字节)。"""

    from app.document.docx_redline import export_redline

    folder = Path(tempfile.mkdtemp(prefix="pra_redline_"))
    target = folder / "revised.docx"
    result = export_redline(data, rounds, str(target), **kwargs)
    return result, target.read_bytes()


def part(data: bytes, name: str) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read(name).decode("utf-8")


@unittest.skipUnless(HAS_DOCX, "未安装 python-docx / lxml")
class TestDocxReading(unittest.TestCase):
    def setUp(self) -> None:
        self.data = build_paper()

    def test_text_matches_loader_byte_for_byte(self) -> None:
        """读出来的文字必须与加载器完全一致，否则字符偏移全部错位。"""

        source = DocxSource(self.data)
        loaded = load_bytes(self.data, filename="paper.docx")
        self.assertEqual(source.text, loaded.text)
        self.assertTrue(loaded.can_write_back)

    def test_headings_formula_image_table_are_all_visible(self) -> None:
        text = DocxSource(self.data).text
        self.assertIn("# 面向小样本场景的图神经网络方法", text)
        self.assertIn("【公式：", text)
        self.assertIn("图 3.1 模型结构", text)
        self.assertIn("| 方法 | 准确率 |", text)
        self.assertIn("85.2", text)

    def test_code_block_keeps_line_breaks(self) -> None:
        """代码段里的软回车要变成换行，否则代码会糊成一行。"""

        text = DocxSource(self.data).text
        self.assertIn("range(n):\n", text)
        self.assertIn("total += weight[i]", text)

    def test_blocks_know_what_they_cannot_edit(self) -> None:
        source = DocxSource(self.data)
        formula = next(block for block in source.blocks if "math" in block.flags)
        picture = next(block for block in source.blocks if "object" in block.flags)
        self.assertTrue(formula.pieces)
        self.assertTrue(picture.obstacles)

    def test_every_span_maps_into_a_block(self) -> None:
        """段落解析出来的每一个块，都要能落到 Word 的某个段落或表格上。"""

        source = DocxSource(self.data)
        spans = build_spans(source.text)
        self.assertTrue(spans)
        for span in spans:
            block = source.block_at(span.start)
            self.assertIsNotNone(block, f"{span.label} 落不到任何块上")
            self.assertGreaterEqual(span.start, block.start)
            # 块的 end 是"下一行开始"，标题这类块会多带一个换行
            self.assertLessEqual(span.end, block.end + 1)
            self.assertIn(span.text, block.text)


@unittest.skipUnless(HAS_DOCX, "未安装 python-docx / lxml")
class TestRedline(unittest.TestCase):
    def setUp(self) -> None:
        self.data = build_paper()
        self.text = DocxSource(self.data).text

    def test_old_applied_code_patch_is_blocked_at_word_export(self) -> None:
        document = Document()
        document.add_paragraph(
            "query_data, query_labels = fs_dataset.load_query_set(FEWSHOT_QUERY_DIR)\n"
            "known_count = len(query_data)"
        )
        data = io.BytesIO()
        document.save(data)
        source_text = DocxSource(data.getvalue()).text
        span = next(s for s in build_spans(source_text) if "FEWSHOT_QUERY_DIR" in s.text)
        start = span.text.index("FEWSHOT_QUERY_DIR")
        stale_patch = Patch(
            "patch-older",
            span.id,
            span_label=span.label,
            quote="FEWSHOT_QUERY_DIR",
            before="FEWSHOT_QUERY_DIR",
            replace="fewshot_query_dir",
            status=STATUS_APPLIED,
            start_in_span=start,
            end_in_span=start + len("FEWSHOT_QUERY_DIR"),
        )
        result, output = export(data.getvalue(), [[stale_patch.to_json()]])
        self.assertEqual(result.applied, [])
        self.assertEqual(len(result.skipped), 1)
        self.assertIn("代码内容或变量名受到保护", result.skipped[0]["reason"])
        self.assertEqual(DocxSource(output).text, source_text)
        self.assertNotIn("fewshot_query_dir", part(output, "word/document.xml"))

    def test_simple_replacement_becomes_tracked_change(self) -> None:
        new_text, rounds = applied_round(
            self.text,
            label="P01",
            quote="在三个公开数据集上取得了更好的效果",
            replace="在 Cora、Citeseer 与 Pubmed 三个数据集上取得了更好的效果",
        )
        result, out_bytes = export(self.data, [rounds])

        self.assertEqual(len(result.applied), 1)
        self.assertEqual(result.skipped, [])
        self.assertEqual(result.text, new_text)

        document_xml = part(out_bytes, "word/document.xml")
        self.assertIn("<w:ins", document_xml)
        self.assertIn("<w:del", document_xml)
        self.assertIn("<w:delText", document_xml)
        self.assertIn("论文修改助手", document_xml)
        self.assertIn("Cora、Citeseer", document_xml)
        self.assertIn("<w:trackRevisions/>", part(out_bytes, "word/settings.xml"))

        # 批注：正文里有三件套锚点，comments.xml 里有批注正文
        self.assertIn("让表述更具体", part(out_bytes, "word/comments.xml"))
        self.assertIn("commentRangeStart", document_xml)
        self.assertIn("commentRangeEnd", document_xml)
        self.assertIn("commentReference", document_xml)

        # 回写之后的文字 == 用户点"全部接受"后看到的样子 == 内存里的结果
        self.assertEqual(DocxSource(out_bytes).text, new_text)

    def test_untouched_part_is_byte_identical(self) -> None:
        new_text, rounds = applied_round(
            self.text,
            label="P01",
            quote="取得了更好的效果",
            replace="取得了更好的效果（在 5 个数据集上验证）",
        )
        result, _ = export(self.data, [rounds])

        expected = self.text.replace("取得了更好的效果", "取得了更好的效果（在 5 个数据集上验证）")
        self.assertEqual(result.text, expected)
        self.assertEqual(result.text, new_text)

    def test_whole_paragraph_rewrite_cannot_destroy_existing_formula(self) -> None:
        """A plain-text replacement must never erase an existing OMML object."""

        spans = build_spans(self.text)
        formula_span = next(span for span in spans if "【公式：" in span.text)
        patch = Patch(
            patch_id="patch-001",
            span_id=formula_span.id,
            span_label=formula_span.label,
            quote=formula_span.text,
            replace="损失函数采用交叉熵，见式（3-1）。",
            rationale="公式排版无法写回时给出说明。",
        )
        patch = validate_patch(patch, spans)
        self.assertEqual(patch.status, STATUS_PROPOSED)
        decide_patch(patch, "accept")
        outcome = apply_patches(self.text, spans, [patch])
        rounds = [item.to_json() for item in outcome.applied]

        result, out_bytes = export(self.data, [rounds])
        self.assertEqual(len(result.applied), 0)
        self.assertEqual(len(result.skipped), 1)
        self.assertEqual(result.text, self.text)
        self.assertIn("公式受到保护", result.skipped[0]["reason"])

    def test_partial_change_reaching_into_formula_is_refused(self) -> None:
        """
        只改公式前面那半句、却把公式也圈进替换范围时，
        回写层必须拒绝（而不是把公式删掉）。
        """

        spans = build_spans(self.text)
        formula_span = next(span for span in spans if "【公式：" in span.text)
        # 覆盖"公式前半段文字 + 整个公式"，但不是整段
        quote = formula_span.text[len("损失函数") :]
        patch = Patch(
            patch_id="patch-001",
            span_id=formula_span.id,
            span_label=formula_span.label,
            quote=quote,
            replace="定义为交叉熵损失。",
            rationale="测试跨公式的改动。",
        )
        patch = validate_patch(patch, spans)
        self.assertEqual(patch.status, STATUS_PROPOSED, patch.conflict_reason)
        decide_patch(patch, "accept")
        outcome = apply_patches(self.text, spans, [patch])
        rounds = [item.to_json() for item in outcome.applied]
        self.assertTrue(rounds)

        result, out_bytes = export(self.data, [rounds])
        self.assertEqual(result.applied, [])
        self.assertEqual(len(result.skipped), 1)
        self.assertIn("公式", result.skipped[0]["reason"])
        # 写不回的改动会在原地留一条"能照着做"的批注：
        # 这个问题需要人工修改，因为…… / 建议修改为：…… / 原因：……
        comments = part(out_bytes, "word/comments.xml")
        self.assertIn("无法修改，需人工修改", comments)
        self.assertIn("修改建议：", comments)
        self.assertIn("原因：", comments)
        self.assertIn("定义为交叉熵损失", comments)  # 建议改成什么
        # 摘要里也能看到"哪一条、在哪、原文与建议"
        summary = result.to_json()
        self.assertEqual(summary["skipped"], 1)
        item = summary["skipped_items"][0]
        self.assertTrue(item["before"])
        self.assertEqual(item["after"], "定义为交叉熵损失。")
        self.assertTrue(item["position"])
        # 原文一个字没动
        self.assertTrue(DocxSource(out_bytes).text.startswith(self.text.split("\n\n")[0]))
        self.assertIn("【公式：", DocxSource(out_bytes).text)

    def test_table_cell_change_is_applied(self) -> None:
        new_text, rounds = applied_round(self.text, label="T01", quote="85.2", replace="87.4")
        result, out_bytes = export(self.data, [rounds])

        self.assertEqual(len(result.applied), 1, result.skipped)
        self.assertEqual(result.text, new_text)
        self.assertIn("| 方法 | 准确率 |", result.text)
        self.assertIn("| 本文方法 | 87.4 |", result.text)
        self.assertIn("87.4", part(out_bytes, "word/document.xml"))

    def test_change_across_table_cells_is_refused(self) -> None:
        """跨单元格的改动（Markdown 竖线在 Word 里不存在）必须拒绝。"""

        spans = build_spans(self.text)
        table_span = next(span for span in spans if span.label == "T01")
        quote = "本文方法 | 85.2"
        patch = Patch(
            patch_id="patch-001",
            span_id=table_span.id,
            span_label=table_span.label,
            quote=quote,
            replace="本文方法 | 87.4",
            rationale="测试跨单元格的改动。",
        )
        patch = validate_patch(patch, spans)
        self.assertEqual(patch.status, STATUS_PROPOSED, patch.conflict_reason)
        decide_patch(patch, "accept")
        outcome = apply_patches(self.text, spans, [patch])

        result, _ = export(self.data, [[item.to_json() for item in outcome.applied]])
        self.assertEqual(result.applied, [])
        self.assertEqual(len(result.skipped), 1)
        self.assertIn("单元格", result.skipped[0]["reason"])

    def test_two_rounds_replay_in_order(self) -> None:
        """多轮会话：第二轮要建立在第一轮的结果上。"""

        after_first, round_one = applied_round(
            self.text, label="P01", quote="更好的效果", replace="更稳定的效果"
        )
        after_second, round_two = applied_round(
            after_first, label="P01", quote="更稳定的效果", replace="更稳定的精度提升"
        )

        result, _ = export(self.data, [round_one, round_two])
        self.assertEqual(len(result.applied), 2, result.skipped)
        self.assertEqual(result.text, after_second)

    def test_replacement_comment_is_one_original_to_final_statement(self):
        from lxml import etree

        from app.document.docx_model import w

        for paragraph, intermediate, final in (
            ("本文采用 MoCo 方法。", "Moco", "MOCO"),
            ("MoCo", "Moco", "MOCO"),
            ("本文采用 MoCo 方法。", "Moco", "MoCo"),
            ("MoCo", "Moco", "MoCo"),
        ):
            with self.subTest(paragraph=paragraph, final=final):
                document = Document()
                document.add_paragraph(paragraph)
                raw = io.BytesIO()
                document.save(raw)
                text = DocxSource(raw.getvalue()).text
                first_text, first = applied_round(
                    text, label="P01", quote="MoCo", replace=intermediate
                )
                final_text, second = applied_round(
                    first_text, label="P01", quote=intermediate, replace=final
                )
                outcome, output = export(raw.getvalue(), [first, second])
                self.assertEqual(outcome.skipped, [])
                self.assertEqual(outcome.text, final_text)
                root = etree.fromstring(part(output, "word/document.xml").encode())
                comments = etree.fromstring(part(output, "word/comments.xml").encode())
                expected = 0 if final == "MoCo" else 1
                self.assertEqual(len(list(comments.iter(w("comment")))), expected)
                self.assertEqual(len(list(root.iter(w("del")))), expected)
                self.assertEqual(len(list(root.iter(w("ins")))), expected)
                if expected:
                    self.assertIn("将「MoCo」改为「MOCO」。", "".join(comments.itertext()))
                    self.assertNotIn("Moco", "".join(comments.itertext()))
                # Reject ALL parts of the replacement restores original;
                # accepting all gives final. Deleting comments is not rejection.
                for accept in (False, True):
                    projected = etree.fromstring(etree.tostring(root))
                    for node in list(projected.iter()):
                        if node.tag not in {w("del"), w("ins")}:
                            continue
                        parent = node.getparent()
                        keep = (node.tag == w("ins")) == accept
                        if keep:
                            index = parent.index(node)
                            for child in list(node):
                                parent.insert(index, child)
                                index += 1
                        parent.remove(node)
                    content = "".join(
                        n.text or "" for n in projected.iter() if n.tag in {w("t"), w("delText")}
                    )
                    self.assertEqual(content, (final_text if accept else text).rstrip("\n"))

    def test_full_rewrite_after_partial_edit_has_no_nested_revisions(self):
        from lxml import etree

        from app.document.docx_model import w

        first_text, first = applied_round(
            self.text, label="P01", quote="更好的效果", replace="更稳定的效果"
        )
        replacement = "本文方法在三个公开数据集上取得稳定提升。"
        _, second = applied_round(
            first_text,
            label="P01",
            quote=BODY.replace("更好的效果", "更稳定的效果"),
            replace=replacement,
        )
        result, output = export(self.data, [first, second])
        self.assertEqual(result.skipped, [])
        root = etree.fromstring(part(output, "word/document.xml").encode())
        self.assertEqual(len(list(root.iter(w("del")))), 1)
        self.assertEqual(len(list(root.iter(w("ins")))), 1)
        self.assertEqual(result.comments, 1)
        self.assertIn(f"将「{BODY}」改为「{replacement}」。", part(output, "word/comments.xml"))

    def test_consolidation_preserves_preexisting_comments_and_anchor_ids(self):
        from lxml import etree

        from app.document.docx_model import w
        from app.document.docx_redline import RedlineWriter

        source = DocxSource(self.data)
        writer = RedlineWriter(source)
        paragraph = next(
            p for p in source.document_root.iter(w("p")) if BODY in "".join(p.itertext(w("t")))
        )
        writer.note_on_paragraph(paragraph, "作者原有批注，不应被删除")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "original.docx"
            source.save(str(path))
            original = path.read_bytes()
        text, first = applied_round(
            self.text, label="P01", quote="更好的效果", replace="更稳定的效果"
        )
        _, second = applied_round(
            text,
            label="P01",
            quote=BODY.replace("更好的效果", "更稳定的效果"),
            replace="本文方法在三个公开数据集上取得稳定提升。",
        )
        result, output = export(original, [first, second])
        self.assertEqual(result.comments, 1)
        root = etree.fromstring(part(output, "word/document.xml").encode())
        comments = etree.fromstring(part(output, "word/comments.xml").encode())
        self.assertEqual(len(list(comments.iter(w("comment")))), 2)
        self.assertIn("作者原有批注，不应被删除", "".join(comments.itertext()))
        ids = {n.get(w("id")) for n in comments.iter(w("comment"))}
        for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
            anchors = [n.get(w("id")) for n in root.iter(w(tag))]
            self.assertEqual(set(anchors), ids)
            self.assertEqual(len(anchors), len(ids))

    def test_round_trip_without_comments_also_removes_noop_revisions(self):
        text, first = applied_round(
            self.text, label="P01", quote="更好的效果", replace="更稳定的效果"
        )
        _, second = applied_round(text, label="P01", quote="更稳定的效果", replace="更好的效果")
        result, output = export(self.data, [first, second], comments=False)
        self.assertEqual(result.text, self.text)
        self.assertEqual(result.comments, 0)
        self.assertNotIn("<w:ins", part(output, "word/document.xml"))
        self.assertNotIn("<w:del", part(output, "word/document.xml"))

    def test_many_changes_in_one_paragraph(self) -> None:
        """
        同一段里的多处改动：Word 里一段常常只有一个 run，
        必须先按改动边界切开，否则第一处就会把后面几处的目标拆掉。
        """

        spans = build_spans(self.text)
        quotes = ["本文", "提出", "方法", "三个", "公开数据集", "更好的效果"]
        patches = []
        for index, quote in enumerate(quotes):
            patch = Patch(
                patch_id=f"patch-{index:03d}",
                span_id="",
                span_label="P01",
                quote=quote,
                replace=quote + "X",
                rationale=f"第 {index + 1} 处。",
            )
            patch = validate_patch(patch, spans)
            self.assertEqual(patch.status, STATUS_PROPOSED, patch.conflict_reason)
            decide_patch(patch, "accept")
            patches.append(patch)

        outcome = apply_patches(self.text, spans, patches)
        result, _ = export(self.data, [[item.to_json() for item in outcome.applied]])

        self.assertEqual(len(result.applied), len(quotes), result.skipped)
        self.assertEqual(result.skipped, [])
        self.assertEqual(result.text, outcome.document)

    def test_change_spanning_several_runs(self) -> None:
        """被 Word 拆成很多个 run 的文字，也要能整段替换。"""

        document = Document()
        paragraph = document.add_paragraph()
        for piece in [
            "本",
            "文提",
            "出的方",
            "法在",
            "三个",
            "公开",
            "数据集上",
            "取得了",
            "更好的",
            "效果。",
        ]:
            paragraph.add_run(piece)
        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()

        text = DocxSource(data).text
        new_text, rounds = applied_round(
            text, label="P01", quote="三个公开数据集上", replace="五个公开数据集上"
        )
        result, out_bytes = export(data, [rounds])

        self.assertEqual(len(result.applied), 1, result.skipped)
        self.assertEqual(result.text, new_text)
        document_xml = part(out_bytes, "word/document.xml")
        self.assertIn("三个", document_xml)
        self.assertIn("五个", document_xml)

    def test_image_is_never_touched(self) -> None:
        """图片段落里的图片必须原样保留（只改图注文字）。"""

        new_text, rounds = applied_round(
            self.text, label="P04", quote="图 3.1 模型结构：", replace="图 3-1 模型结构："
        )
        result, out_bytes = export(self.data, [rounds])
        self.assertEqual(len(result.applied), 1, result.skipped)
        self.assertEqual(result.text, new_text)
        with zipfile.ZipFile(io.BytesIO(out_bytes)) as archive:
            self.assertTrue(any(name.startswith("word/media/") for name in archive.namelist()))

    def test_pure_deletion_is_tracked(self) -> None:
        """删除类改动：只留删除痕迹，不插入任何文字。"""

        spans = build_spans(self.text)
        patch = Patch(
            patch_id="patch-001",
            span_id="",
            span_label="P02",
            quote="为了验证有效性，",
            replace="",
            delete=True,
            rationale="这句话是套话，删掉。",
        )
        patch = validate_patch(patch, spans)
        self.assertEqual(patch.status, STATUS_PROPOSED, patch.conflict_reason)
        decide_patch(patch, "accept")
        outcome = apply_patches(self.text, spans, [patch])
        rounds = [item.to_json() for item in outcome.applied]

        result, out_bytes = export(self.data, [rounds])
        self.assertEqual(len(result.applied), 1, result.skipped)
        self.assertEqual(result.text, outcome.document)
        self.assertNotIn("为了验证有效性", result.text)
        self.assertIn(
            "<w:delText>为了验证有效性，</w:delText>", part(out_bytes, "word/document.xml")
        )

    def test_document_without_settings_part_is_wired_up(self) -> None:
        """
        原文件没有 settings.xml 时（有些导出器就是不给），
        要自己补上 settings 部件、关系与内容类型，否则 Word 会报文件损坏。
        """

        buffer = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(self.data)) as source:
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
                for name in source.namelist():
                    if name == "word/settings.xml":
                        continue
                    target.writestr(name, source.read(name))
        stripped = buffer.getvalue()

        _, rounds = applied_round(
            self.text, label="P01", quote="更好的效果", replace="更明显的效果"
        )
        result, out_bytes = export(stripped, [rounds])
        self.assertEqual(len(result.applied), 1, result.skipped)
        self.assertIn("<w:trackRevisions/>", part(out_bytes, "word/settings.xml"))
        self.assertIn("settings.xml", part(out_bytes, "word/_rels/document.xml.rels"))
        self.assertIn("/word/settings.xml", part(out_bytes, "[Content_Types].xml"))

    def test_output_reopens_as_a_normal_word_file(self) -> None:
        """产物必须是一份"正常"的 docx：能被常规库重新打开并读到批注与痕迹。"""

        new_text, rounds = applied_round(
            self.text, label="P01", quote="更好的效果", replace="更明显的效果"
        )
        result, out_bytes = export(self.data, [rounds])
        self.assertEqual(result.text, new_text)

        import tempfile
        from pathlib import Path

        folder = Path(tempfile.mkdtemp(prefix="pra_reopen_"))
        target = folder / "revised.docx"
        target.write_bytes(out_bytes)
        reopened = Document(str(target))
        paragraphs = [paragraph.text for paragraph in reopened.paragraphs]
        # python-docx 的 paragraph.text 只认段落直属的 <w:r>：
        # 被删掉的字在 <w:del> 里、新插入的字在 <w:ins> 里，它都看不到。
        # 所以这里应该正好剩下"改动范围之外"的那两截文字。
        self.assertIn("本文提出的方法在三个公开数据集上取得了。", paragraphs)
        self.assertTrue(any("为了验证有效性" in item for item in paragraphs))

        rels = part(out_bytes, "word/_rels/document.xml.rels")
        self.assertIn("comments.xml", rels)
        self.assertIn("/word/comments.xml", part(out_bytes, "[Content_Types].xml"))
        # 每个改动配一条批注
        document_xml = part(out_bytes, "word/document.xml")
        self.assertEqual(
            document_xml.count("<w:commentReference"),
            len(result.applied) + (1 if result.skipped else 0),
        )


class TestReplay(unittest.TestCase):
    """回写靠"重放"：必须按轮次、按顺序取"真正落实"的改动（不需要 Word 环境）。"""

    def test_report_wiring_writes_word_copy(self) -> None:
        """
        接线检查：报告节点落盘时，要顺手把改动写回 Word（另存副本），
        并把路径与摘要写回 state。
        """

        if not HAS_DOCX:
            self.skipTest("未安装 python-docx / lxml")
        try:
            from app.agents.report import write_redline_docx
        except ImportError as exc:  # pragma: no cover - 取决于运行环境
            self.skipTest(f"报告模块不可用：{exc}")

        data = build_paper()
        folder = Path(tempfile.mkdtemp(prefix="pra_report_"))
        source = folder / "paper.docx"
        source.write_bytes(data)
        text = DocxSource(data).text
        new_text, rounds = applied_round(
            text, label="P01", quote="更好的效果", replace="更明显的效果"
        )

        state = {
            "source_docx": str(source),
            "revision_history": [{"stage": "apply", "applied": rounds}],
        }
        update = write_redline_docx(state, run_dir=folder, logger=lambda message: None)

        self.assertTrue(update.get("docx_path"))
        produced = Path(update["docx_path"])
        self.assertTrue(produced.exists())
        self.assertEqual(update["redline_summary"]["applied"], 1)
        self.assertEqual(DocxSource(produced.read_bytes()).text, new_text)
        # 原文件必须一字不动
        self.assertEqual(source.read_bytes(), data)

    def test_no_source_file_means_no_word_copy(self) -> None:
        try:
            from app.agents.report import write_redline_docx
        except ImportError as exc:  # pragma: no cover
            self.skipTest(f"报告模块不可用：{exc}")

        folder = Path(tempfile.mkdtemp(prefix="pra_report_"))
        update = write_redline_docx(
            {"revision_history": [{"stage": "apply", "applied": [{"patch_id": "a"}]}]},
            run_dir=folder,
            logger=lambda message: None,
        )
        self.assertEqual(update, {})

    def test_applied_rounds_keeps_only_real_apply_steps(self) -> None:
        from app.document.docx_redline import applied_rounds

        state = {
            "revision_history": [
                {"stage": "apply", "applied": [{"patch_id": "a"}]},
                {"stage": "revision", "applied": [{"patch_id": "x"}]},
                {"stage": "apply", "applied": []},
                {"stage": "apply", "applied": [{"patch_id": "b"}]},
            ]
        }
        rounds = applied_rounds(state)
        self.assertEqual(
            [[item["patch_id"] for item in group] for group in rounds],
            [["a"], ["b"]],
        )

    def test_applied_rounds_is_empty_without_history(self) -> None:
        from app.document.docx_redline import applied_rounds

        self.assertEqual(applied_rounds({}), [])


def add_ref_field(paragraph, instruction: str, result: str) -> None:
    """
    在段落里插入一个 Word 域（begin / instrText / separate / 结果 / end）。

    完全照搬论文里引文序号的样子——[4] 这类不是打上去的字符，
    而是 REF 交叉引用域，改它要同步后面的参考文献。
    """

    from docx.oxml.ns import qn

    def fld(char_type: str):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), char_type)
        paragraph.add_run()._element.append(element)

    fld("begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    paragraph.add_run()._element.append(instr)
    fld("separate")
    paragraph.add_run(result)
    fld("end")


def docx_with_citation_field() -> bytes:
    """正文里带一个引文域：Transformer[4]、GNN——[4] 是 REF 域。"""

    document = Document()
    document.add_paragraph("近年的研究。")
    paragraph = document.add_paragraph()
    paragraph.add_run("研究者们尝试运用Transformer")
    add_ref_field(paragraph, " REF _Ref25657 \\r \\h ", "[4]")
    paragraph.add_run("、GNN等网络结构，用于提取语义特征。")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_with_blank_line_run() -> bytes:
    """
    构造一份"归一化会删字符"的文档。

    代码段里有三个连续换行（``a<break><break><break>b``）。
    整篇归一化会把 ``\\n{3,}`` 压成 ``\\n\\n``——**删掉一个字符**，
    于是它后面所有块的偏移整体左移。这正是线上那份 217 个块错位的成因。
    """

    document = Document()
    document.add_heading("1 引言", level=2)
    document.add_paragraph("第一段文字。")

    code = document.add_paragraph()
    code.add_run("a")
    code.add_run().add_break()
    code.add_run().add_break()
    code.add_run().add_break()
    code.add_run("b")

    document.add_heading("2 方法", level=2)
    document.add_paragraph("第二段文字。")
    document.add_paragraph("第三段文字。")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@unittest.skipUnless(HAS_DOCX, "未安装 python-docx / lxml")
class TestBlockAlignment(unittest.TestCase):
    """
    块的偏移必须与扁平文本自洽。

    历史 bug：``block.start/end`` 按**归一化前**的长度累加，而 ``self.text``
    是归一化**之后**的文本。归一化会删字符（连续空行、行尾空白、首尾空白），
    删了之后后面所有块的偏移整体左移，
    ``block_at(span.start)`` 就会拿到前一个块，让回写层误报
    "这条改动跨了文档结构，需要人工处理"——而补丁本身明明只是句内改字。
    """

    def setUp(self) -> None:
        self.data = docx_with_blank_line_run()
        self.source = DocxSource(self.data)

    def test_text_is_unchanged_by_the_fix(self) -> None:
        """安全底线：归一化改成逐块做之后，模型看到的文本必须逐字节不变。"""

        from app.utils.text import normalize_document_text

        raw = [block.text for block in self.source._collect_blocks()]
        self.assertEqual(
            self.source.text,
            normalize_document_text("\n\n".join(raw)),
        )

    def test_fixture_really_shrinks_under_normalization(self) -> None:
        """
        确认这份样本**真的**会触发"归一化删字符"。

        没有这条，上面的对齐测试可能在原地空转：
        如果样本本身不触发删除，旧代码也不会错位，测试就抓不到这个 bug。
        """

        from app.utils.text import normalize_document_text

        raw = [block.text for block in self.source._collect_blocks()]
        joined = "\n\n".join(raw)
        # normalize_document_text 正常会补一个结尾换行；被删过字符就会小于 len+1
        self.assertLess(
            len(normalize_document_text(joined)),
            len(joined) + 1,
            "样本没有触发归一化删字符，测试是空转",
        )

    def test_every_block_offset_matches_its_own_text(self) -> None:
        for block in self.source.blocks:
            self.assertEqual(
                self.source.text[block.start : block.start + len(block.text)],
                block.text,
                f"块 {block.index} 的偏移与文本对不上",
            )

    def test_no_misaligned_block_warning(self) -> None:
        self.assertFalse(
            [w for w in self.source.warnings if "无法与原文对齐" in w],
            self.source.warnings,
        )

    def test_pieces_still_point_at_the_same_text(self) -> None:
        """每个片段的偏移搬过去之后，取出来的还得是它原来那段文字。"""

        for block in self.source.blocks:
            for piece in block.pieces:
                self.assertEqual(
                    block.text[piece.start : piece.end],
                    piece.rendered,
                    f"块 {block.index} 的片段偏移错位",
                )

    def test_patch_after_the_shifted_part_is_still_writable(self) -> None:
        """端到端：在漂移点**之后**的段落上，普通正文补丁必须还能写回 Word。"""

        spans = build_spans(self.source.text)
        target = next(span for span in spans if "第三段文字" in span.text)

        _new_text, rounds = applied_round(
            self.source.text,
            label=target.label,
            quote="第三段文字",
            replace="第三段内容",
        )
        result, out_bytes = export(self.data, [rounds])

        self.assertEqual(result.skipped, [], "补丁被误判为无法写回")
        self.assertEqual(len(result.applied), 1)
        self.assertIn("第三段内容", DocxSource(out_bytes).text)


@unittest.skipUnless(HAS_DOCX, "未安装 python-docx / lxml")
class TestCitationFields(unittest.TestCase):
    """
    引文序号（[4]、[5]）是 Word 交叉引用域，不是普通文字。

    规则：**改动范围只要碰到引文域，一律不改**，改成报告，
    原因写清楚"需要同步后面文献引用"——因为改引文要连带改参考文献编号，
    机器判断不了，只能交给人。
    """

    def setUp(self) -> None:
        self.data = docx_with_citation_field()
        self.text = DocxSource(self.data).text

    def test_field_result_is_read_as_plain_text(self) -> None:
        self.assertIn("Transformer[4]、GNN等网络结构", self.text)

    def test_change_spanning_the_field_is_refused_with_the_citation_reason(self) -> None:
        spans = build_spans(self.text)
        target = next(span for span in spans if "Transformer[4]" in span.text)

        _new_text, rounds = applied_round(
            self.text,
            label=target.label,
            quote="Transformer[4]、GNN等",
            replace="Transformer[4]、KNN等",
        )
        result, _out = export(self.data, [rounds])

        self.assertEqual(result.applied, [])
        self.assertEqual(len(result.skipped), 1)
        self.assertIn("这条改动需要同步后面文献引用", result.skipped[0]["reason"])

    def test_change_before_the_field_is_still_writable(self) -> None:
        """不碰域的改动照旧能写回——不能因为文档里有引文就整体摆烂。"""

        spans = build_spans(self.text)
        target = next(span for span in spans if "近年的研究" in span.text)

        _new_text, rounds = applied_round(
            self.text,
            label=target.label,
            quote="近年的研究",
            replace="近年来的研究",
        )
        result, out_bytes = export(self.data, [rounds])

        self.assertEqual(result.skipped, [])
        self.assertIn("近年来的研究", DocxSource(out_bytes).text)


@unittest.skipUnless(HAS_DOCX, "未安装 python-docx / lxml")
class TestOriginalRevisionBoundaries(unittest.TestCase):
    def test_term_across_old_insertion_keeps_unedited_revision(self):
        from app.document.docx_model import w

        document = Document()
        paragraph = document.add_paragraph("方法 Mo")
        insertion = OxmlElement("w:ins")
        insertion.set(w("author"), "原作者")
        insertion.set(w("id"), "99")
        run = OxmlElement("w:r")
        node = OxmlElement("w:t")
        node.text = "CO 和旧修订保留"
        run.append(node)
        insertion.append(run)
        paragraph._p.append(insertion)
        paragraph.add_run("。")
        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()
        text = DocxSource(data).text
        _, changes = applied_round(
            text, label=build_spans(text)[0].label, quote="MoCO", replace="MoCo"
        )
        result, output = export(data, [changes])
        self.assertEqual(result.skipped, [])
        self.assertEqual(DocxSource(output).text, text.replace("MoCO", "MoCo"))
        self.assertEqual(result.comments, 1)
        root = DocxSource(output).document_root
        old = [n for n in root.iter(w("ins")) if n.get(w("author")) == "原作者"]
        self.assertEqual(
            "".join(n.text or "" for x in old for n in x.iter(w("t"))), " 和旧修订保留"
        )
        self.assertIn("MoCO", "".join(n.text or "" for n in root.iter(w("delText"))))

    def test_several_text_nodes_in_one_run_do_not_lose_neighbours(self):
        document = Document()
        paragraph = document.add_paragraph()
        run = paragraph.add_run()._r
        for value in ["原文 ", "MoCO", " 和 ", "batch", " 保留尾部"]:
            node = OxmlElement("w:t")
            node.text = value
            run.append(node)
        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()
        text = DocxSource(data).text
        label = build_spans(text)[0].label
        _, one = applied_round(text, label=label, quote="MoCO", replace="MoCo")
        _, two = applied_round(text, label=label, quote="batch", replace="批次")
        two[0]["patch_id"] = "patch-002"
        result, output = export(data, [one + two])
        self.assertEqual(result.skipped, [])
        self.assertEqual(
            DocxSource(output).text, text.replace("MoCO", "MoCo").replace("batch", "批次")
        )
        self.assertEqual(result.comments, 2)

    def test_one_term_spans_multiple_text_nodes_in_same_run(self):
        document = Document()
        paragraph = document.add_paragraph("方法 ")
        run = paragraph.add_run()._r
        for value in ["Mo", "CO", " 保留"]:
            node = OxmlElement("w:t")
            node.text = value
            run.append(node)
        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()
        text = DocxSource(data).text
        _, changes = applied_round(
            text, label=build_spans(text)[0].label, quote="MoCO", replace="MoCo"
        )
        result, output = export(data, [changes])
        self.assertEqual(result.skipped, [])
        self.assertEqual(DocxSource(output).text, text.replace("MoCO", "MoCo"))


if __name__ == "__main__":
    unittest.main()
